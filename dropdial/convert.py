"""Format conversions. Every function takes a source path and returns a list of
output paths written next to the source."""
from __future__ import annotations

import html
import io
import re
import sys
from pathlib import Path

from PIL import Image, ImageOps

from .util import UserError, sibling

try:  # HEIC / HEIF support
    import pillow_heif
    pillow_heif.register_heif_opener()
    HEIF_OK = True
except Exception:  # pragma: no cover
    HEIF_OK = False

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore
try:
    # colour-manage PDFs: photos and CMYK artwork inside PDFs keep their real colours
    # (off by default in some PyMuPDF versions, which badly shifts them)
    fitz.TOOLS.set_icc(True)
except Exception:  # pragma: no cover
    pass

# most-used first: the dial shows the first 8, the rest go on its "More formats" side
IMAGE_TARGETS = ["PNG", "JPG", "WEBP", "PDF", "HEIC", "AVIF", "TIFF", "GIF", "BMP", "ICO", "SVG", "ICNS",
                 "TGA", "QOI"]
EXTRA_TARGETS = {"ICNS", "TGA", "QOI"}       # written by formats/rawimages.save_extra
PDF_TARGETS = ["JPG", "PNG", "DOCX", "TXT"]
DOCX_TARGETS = ["PDF", "TXT"]
TEXT_TARGETS = ["PDF", "DOCX"]
VECTOR_TARGETS = ["PNG", "JPG", "WEBP", "PDF", "ICO"]
EBOOK_TARGETS = ["PDF", "DOCX", "TXT"]

PIL_FORMAT = {"PNG": "PNG", "JPG": "JPEG", "WEBP": "WEBP", "HEIC": "HEIF", "TIFF": "TIFF",
              "AVIF": "AVIF", "BMP": "BMP", "PDF": "PDF", "ICO": "ICO", "GIF": "GIF"}
NO_ALPHA = {"JPG", "PDF", "BMP"}


# ---------------------------------------------------------------- images
_SRGB = None


def _srgb_profile():
    global _SRGB
    if _SRGB is None:
        from PIL import ImageCms
        _SRGB = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB"))
    return _SRGB


def _to_8bit(im: Image.Image) -> Image.Image:
    """16-bit and 32-bit greyscale (I;16, I, F) -> 8-bit L. A plain convert() clips them to white."""
    import numpy as np
    a = np.asarray(im, dtype=np.float64)
    if im.mode == "F":
        scale = 255.0 if a.max(initial=0) <= 1.0 else 1.0
    else:
        scale = 1 / 257 if a.max(initial=0) > 255 else 1.0
    return Image.fromarray(np.clip(a * scale + 0.5, 0, 255).astype(np.uint8), "L")


def to_srgb(im: Image.Image, icc: bytes | None) -> Image.Image:
    """Convert pixels from the file's own colour profile (Display P3 from iPhones and Macs,
    Adobe RGB, ProPhoto, CMYK print profiles…) to standard sRGB, so every output looks
    like the original no matter which format it's saved in. Returns RGB, RGBA or L."""
    if im.mode in ("I;16", "I;16B", "I;16L", "I;16N", "I", "F"):
        im = _to_8bit(im)
    if im.mode in ("P", "PA", "LA", "La", "RGBa"):
        im = im.convert("RGBA")
    if icc and im.mode in ("RGB", "RGBA", "CMYK"):
        try:
            from PIL import ImageCms
            src = ImageCms.ImageCmsProfile(io.BytesIO(icc))
            alpha = im.getchannel("A") if im.mode == "RGBA" else None
            base = im.convert("RGB") if im.mode == "RGBA" else im
            base = ImageCms.profileToProfile(base, src, _srgb_profile(),
                                             renderingIntent=ImageCms.Intent.RELATIVE_COLORIMETRIC,
                                             outputMode="RGB")   # same as PDF viewers (MuPDF)
            if alpha is not None:
                base = base.convert("RGBA")
                base.putalpha(alpha)
            im = base
        except Exception:
            pass   # a broken or unusual profile: fall back to the raw pixels below
    if im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    im.info.pop("icc_profile", None)   # pixels are sRGB now: never carry the old profile along
    return im


def open_image(path: str | Path) -> Image.Image:
    """Decode an image upright (camera rotation applied) and in sRGB."""
    im = Image.open(path)
    im.seek(0)
    icc = im.info.get("icc_profile")
    im = ImageOps.exif_transpose(im)
    return to_srgb(im, icc)


def open_preview(path: str | Path, max_side: int = 1400) -> Image.Image:
    """Fast, reduced-size decode for on-screen previews (JPEG decodes at 1/2, 1/4, 1/8 scale)."""
    im = Image.open(path)
    im.seek(0)
    icc = im.info.get("icc_profile")
    if im.format == "JPEG":
        im.draft(im.mode if im.mode in ("RGB", "L", "CMYK") else "RGB", (max_side, max_side))
    im = ImageOps.exif_transpose(im)
    im = to_srgb(im, icc)
    if max(im.size) > max_side:
        im.thumbnail((max_side, max_side), Image.LANCZOS)
    return im


def open_pdf(path: str | Path):
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise UserError("This PDF is password-protected. Remove the password first.")
    return doc


def flatten(im: Image.Image, bg=(255, 255, 255)) -> Image.Image:
    if im.mode == "RGBA":
        base = Image.new("RGB", im.size, bg)
        base.paste(im, mask=im.split()[3])
        return base
    return im.convert("RGB") if im.mode != "RGB" else im


def save_image(im: Image.Image, out: Path, target: str, quality: int | None = None, exif=None) -> None:
    from .util import CONFIG
    quality = quality or int(CONFIG.get("jpg_quality", 90))
    target = target.upper()
    fmt = PIL_FORMAT[target]
    if target in NO_ALPHA:
        im = flatten(im)
    kw: dict = {}
    if target in ("JPG", "WEBP", "AVIF", "HEIC"):
        kw["quality"] = quality
    if target == "JPG":
        kw.update(optimize=True, progressive=True)
        if quality >= 90:
            kw["subsampling"] = 0   # 4:4:4: no colour bleeding around text and fine detail
    if target == "PNG":
        kw["optimize"] = True
    if target == "ICO":
        side = min(256, max(im.size))
        im = ImageOps.contain(im.convert("RGBA"), (side, side))
        kw["sizes"] = [(s, s) for s in (16, 24, 32, 48, 64, 128, 256) if s <= side]
    if target == "PDF":
        if not isinstance(out, (str, Path)):
            raise ValueError("PDF output needs a file path")
        write_images_pdf([im], out)
        return
    if exif and target in ("JPG", "WEBP", "PNG", "TIFF"):
        kw["exif"] = exif
    atomic_save(lambda dst: im.save(dst, fmt, **kw), out)


def atomic_save(write, out):
    """Write to a hidden temp file, then rename: Explorer never shows a half-written file."""
    import os
    if not isinstance(out, (str, Path)):
        write(out)  # in-memory buffer
        return
    out = Path(out)
    tmp = out.with_name(f".~{out.name}.part")
    try:
        write(tmp)
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


def image_to(src: str | Path, target: str) -> list[Path]:
    target = target.upper()
    out = sibling(src, ext="." + target.lower())
    if target == "PDF":
        write_images_pdf([src], out)
        return [out]
    if target in EXTRA_TARGETS:
        from .formats.rawimages import save_extra
        save_extra(open_image(src), out, target)
        return [out]
    if target == "SVG":
        svg = trace_svg(open_image(src))
        atomic_save(lambda dst: Path(dst).write_text(svg, encoding="utf-8"), out)
        return [out]
    save_image(open_image(src), out, target)
    return [out]


def can_trace() -> bool:
    try:
        import vtracer  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def trace_svg(im: Image.Image, max_side: int = 2000) -> str:
    """Trace a picture into vector shapes (logos, icons, drawings trace best; photos get large)."""
    import numpy as np
    import vtracer
    im = im.convert("RGBA")
    w, h = im.size
    k = min(1.0, max_side / max(w, h))
    if k < 1:
        im = im.resize((max(1, round(w * k)), max(1, round(h * k))), Image.LANCZOS)
    # reduce to a palette first: the tracer merges smooth gradients into one flat colour
    # otherwise (a sunset became a single pink rectangle). 64 colours keeps photos within ~1/255.
    alpha = im.getchannel("A")
    q = im.convert("RGB").quantize(colors=64, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
    q = q.convert("RGBA")
    q.putalpha(alpha.point(lambda a: 255 if a >= 128 else 0))
    px = [tuple(int(c) for c in row) for row in np.asarray(q).reshape(-1, 4)]
    svg = vtracer.convert_pixels_to_svg(
        px, size=im.size, colormode="color", hierarchical="stacked", mode="spline",
        filter_speckle=4, color_precision=8, layer_difference=1, corner_threshold=60,
        length_threshold=4.0, max_iterations=10, splice_threshold=45, path_precision=2)
    svg = re.sub(r"<!--.*?-->\s*", "", svg, flags=re.S)       # drop the generator comment
    if k < 1:   # keep the original's size: scale the traced shapes back up
        svg = svg.replace(f'width="{im.width}" height="{im.height}"',
                          f'width="{w}" height="{h}" viewBox="0 0 {im.width} {im.height}"', 1)
    return svg


def images_to_single_pdf(srcs: list[str | Path]) -> list[Path]:
    out = sibling(srcs[0], "Combined", ".pdf")
    write_images_pdf(list(srcs), out)
    return [out]


LOSSY_SOURCES = {"JPEG", "MPO", "HEIF", "WEBP", "AVIF"}


def _pdf_image_bytes(item) -> tuple[bytes, tuple[int, int], float]:
    """(encoded image, pixel size, dpi) for one PDF page.

    A plain sRGB JPEG that needs no rotation is embedded byte-for-byte (no quality loss).
    Everything else is converted to sRGB and embedded as PNG (lossless) or, for photos
    that were lossy to begin with, as a high-quality 4:4:4 JPEG."""
    if isinstance(item, Image.Image):
        im, fmt, raw, dpi = item, None, None, 150.0
    else:
        src = Image.open(item)
        fmt = src.format
        dpi = float((src.info.get("dpi") or (0, 0))[0] or 0)
        orient = src.getexif().get(0x0112, 1)
        icc = src.info.get("icc_profile")
        # cameras and screenshots say 72/96/144 dpi, which would make a 12 MP photo a 56-inch
        # page; only trust values that look like a real scan or print
        dpi = dpi if 150 <= dpi <= 1200 else 150.0
        if fmt == "JPEG" and src.mode in ("RGB", "L") and orient in (0, 1) and not icc:
            return Path(item).read_bytes(), src.size, dpi
        raw = item
        im = open_image(item)
    im = flatten(im) if im.mode == "RGBA" else im
    buf = io.BytesIO()
    if fmt in LOSSY_SOURCES:
        im.convert("RGB").save(buf, "JPEG", quality=95, subsampling=0, optimize=True)
    else:
        im.save(buf, "PNG", compress_level=6)
    del raw
    return buf.getvalue(), im.size, dpi


def write_images_pdf(items: list, out: str | Path) -> None:
    """One page per image, each page exactly the image's size at its own DPI (150 if unknown)."""
    doc = fitz.open()
    try:
        for item in items:
            data, (w, h), dpi = _pdf_image_bytes(item)
            k = 72.0 / dpi
            page = doc.new_page(width=w * k, height=h * k)
            page.insert_image(page.rect, stream=data)
        atomic_save(lambda dst: doc.save(str(dst), garbage=3, deflate=True), Path(out))
    finally:
        doc.close()


# ---------------------------------------------------------------- PDF
def pdf_to_images(src: str | Path, target: str, dpi: int | None = None) -> list[Path]:
    from .util import CONFIG
    dpi = dpi or int(CONFIG.get("pdf_dpi", 200))
    target = target.upper()
    outs = []
    with open_pdf(src) as doc:
        n = doc.page_count
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            label = f"page {i + 1}" if n > 1 else ""
            out = sibling(src, label, "." + target.lower())
            save_image(im, out, target)
            outs.append(out)
    return outs


def pdf_to_txt(src: str | Path) -> list[Path]:
    with open_pdf(src) as doc:
        text = "\n\f\n".join(page.get_text("text") for page in doc)
    out = sibling(src, ext=".txt")
    out.write_text(text, encoding="utf-8")
    return [out]


def pdf_to_docx(src: str | Path) -> list[Path]:
    """Text blocks become paragraphs, embedded images are kept, one Word page per PDF page."""
    out = sibling(src, ext=".docx")
    with open_pdf(src) as doc:
        _doc_to_docx(doc, out)
    return [out]


def _doc_to_docx(doc, out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches

    d = Document()
    for pno, page in enumerate(doc):
        if pno:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        items = []
        for b in page.get_text("dict")["blocks"]:
            y = b["bbox"][1]
            if b.get("type") == 0:
                lines = []
                size = 11.0
                bold = False
                for ln in b["lines"]:
                    spans = ln["spans"]
                    if spans:
                        size = max(size, max(s["size"] for s in spans)) if not lines else size
                        bold = bold or any("Bold" in s["font"] for s in spans)
                    lines.append("".join(s["text"] for s in spans))
                txt = " ".join(l.strip() for l in lines).strip()
                if txt:
                    items.append((y, "text", (txt, size, bold)))
            elif b.get("type") == 1 and b.get("image"):
                raw = b["image"]
                if b.get("colorspace") not in (1, 3) or b.get("mask") or b.get("ext") not in ("png", "jpeg", "jpg"):
                    # CMYK, transparency or odd encodings: Word would show them wrong, so use
                    # exactly what the page shows (colour-managed, on white)
                    pix = page.get_pixmap(clip=fitz.Rect(b["bbox"]), dpi=200, alpha=False)
                    raw = pix.tobytes("png")
                items.append((y, "img", (raw, b["bbox"])))
        items.sort(key=lambda t: t[0])
        for _, kind, data in items:
            if kind == "text":
                txt, size, bold = data
                if size >= 16:
                    d.add_heading(txt, level=1 if size >= 20 else 2)
                else:
                    p = d.add_paragraph()
                    r = p.add_run(txt)
                    r.bold = bold and len(txt) < 120
            else:
                raw, bbox = data
                width_in = min(6.5, max(0.5, (bbox[2] - bbox[0]) / 72))
                try:
                    d.add_picture(io.BytesIO(raw), width=Inches(width_in))
                except Exception:
                    pass
    d.save(out)


# ---------------------------------------------------------------- DOCX
def _docx_paragraph_html(p) -> str:
    parts = []
    for r in p.runs:
        t = html.escape(r.text)
        if not t:
            continue
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        if r.underline:
            t = f"<u>{t}</u>"
        parts.append(t)
    body = "".join(parts) or "&nbsp;"
    style = (p.style.name or "").lower() if p.style is not None else ""
    if style.startswith("heading"):
        lvl = "".join(c for c in style if c.isdigit()) or "1"
        lvl = str(min(int(lvl), 6))
        return f"<h{lvl}>{body}</h{lvl}>"
    if style == "title":
        return f"<h1>{body}</h1>"
    if "list" in style:
        return f"<p>&bull; {body}</p>"
    return f"<p>{body}</p>"


def docx_to_html(src: str | Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(src)
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(_docx_paragraph_html(Paragraph(child, d)))
        elif tag == "tbl":
            t = Table(child, d)
            rows = []
            for row in t.rows:
                cells = "".join(f"<td>{html.escape(c.text)}</td>" for c in row.cells)
                rows.append(f"<tr>{cells}</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    return "\n".join(out)


def docx_to_txt(src: str | Path) -> list[Path]:
    from docx import Document
    d = Document(src)
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append("\t".join(c.text for c in row.cells))
    out = sibling(src, ext=".txt")
    out.write_text("\n".join(lines), encoding="utf-8")
    return [out]


def _word_to_pdf(src: Path, out: Path) -> bool:
    """Use Microsoft Word (if installed) for a perfect-fidelity export."""
    if sys.platform != "win32":
        return False
    try:
        import win32com.client  # type: ignore
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        try:
            doc = word.Documents.Open(str(src.resolve()), ReadOnly=True)
            doc.SaveAs(str(out.resolve()), FileFormat=17)  # wdFormatPDF
            doc.Close(False)
        finally:
            word.Quit()
        return out.exists()
    except Exception:
        return False


CSS = """
* { font-family: sans-serif; }
body { font-size: 11pt; line-height: 1.35; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; }
p { margin: 0 0 6pt 0; }
pre { font-family: monospace; font-size: 9.5pt; white-space: pre-wrap; }
table { border-collapse: collapse; margin: 6pt 0; }
td { border: 1px solid #888; padding: 3pt 5pt; font-size: 10pt; }
"""


def html_to_pdf(body_html: str, out: Path) -> None:
    story = fitz.Story(html=body_html, user_css=CSS)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def docx_to_pdf(src: str | Path) -> list[Path]:
    src = Path(src)
    out = sibling(src, ext=".pdf")
    if not _word_to_pdf(src, out):
        html_to_pdf(docx_to_html(src), out)
    return [out]


# ---------------------------------------------------------------- text
def read_text(src: str | Path) -> str:
    raw = Path(src).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def text_to_pdf(src: str | Path) -> list[Path]:
    out = sibling(src, ext=".pdf")
    html_to_pdf(f"<pre>{html.escape(read_text(src))}</pre>", out)
    return [out]


def text_to_docx(src: str | Path) -> list[Path]:
    from docx import Document
    d = Document()
    for line in read_text(src).splitlines():
        d.add_paragraph(line)
    out = sibling(src, ext=".docx")
    d.save(out)
    return [out]


# ---------------------------------------------------------------- SVG
# Qt's SVG renderer handles gradients, percentage sizes and text; MuPDF draws gradients as
# solid black, so it's only the fallback.
_QT_APP = None


def _qt_svg(src: str | Path):
    """A valid QSvgRenderer for the file, or None (Qt SVG unavailable or the file is odd)."""
    global _QT_APP
    try:
        import threading
        from PySide6.QtCore import QByteArray
        from PySide6.QtGui import QGuiApplication
        from PySide6.QtSvg import QSvgRenderer
        if QGuiApplication.instance() is None:
            if threading.current_thread() is not threading.main_thread():
                return None
            import os
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")   # command line / tests
            _QT_APP = QGuiApplication([sys.argv[0] if sys.argv and sys.argv[0] else "dropdial"])
        r = QSvgRenderer(QByteArray(Path(src).read_bytes()))
        return r if r.isValid() else None
    except Exception:
        return None


def _svg_size(r) -> tuple[float, float]:
    ds = r.defaultSize()
    if ds.width() > 0 and ds.height() > 0:
        return float(ds.width()), float(ds.height())
    vb = r.viewBoxF()
    if vb.width() > 0 and vb.height() > 0:
        return vb.width(), vb.height()
    return 512.0, 512.0


def _open_svg(src: str | Path):
    try:
        return fitz.open(str(src))
    except Exception as e:
        raise UserError("This SVG couldn't be read.") from e


def svg_to_image(src: str | Path, long_side: int = 2048) -> Image.Image:
    """Render an SVG to pixels; small icons are scaled up so the result stays crisp."""
    r = _qt_svg(src)
    if r is not None:
        from PySide6.QtCore import QRectF, Qt
        from PySide6.QtGui import QImage, QPainter
        w, h = _svg_size(r)
        zoom = max(1.0, long_side / max(w, h))
        img = QImage(max(1, round(w * zoom)), max(1, round(h * zoom)), QImage.Format_ARGB32_Premultiplied)
        img.fill(Qt.transparent)
        p = QPainter(img)
        p.setRenderHint(QPainter.Antialiasing)
        p.setRenderHint(QPainter.SmoothPixmapTransform)
        r.render(p, QRectF(0, 0, img.width(), img.height()))
        p.end()
        img = img.convertToFormat(QImage.Format_RGBA8888)
        return Image.frombuffer("RGBA", (img.width(), img.height()), bytes(img.constBits()), "raw", "RGBA", 0, 1).copy()
    with _open_svg(src) as doc:
        page = doc[0]
        w, h = page.rect.width or 1, page.rect.height or 1
        zoom = max(1.0, long_side / max(w, h))
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
        mode = "RGBA" if pix.alpha else "RGB"
        return Image.frombytes(mode, (pix.width, pix.height), pix.samples)


def _svg_to_pdf_qt(r, out: Path) -> None:
    from PySide6.QtCore import QMarginsF, QRectF, QSizeF
    from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter
    w, h = _svg_size(r)

    def write(dst):
        pw = QPdfWriter(str(dst))
        pw.setResolution(72)
        pw.setPageLayout(QPageLayout(QPageSize(QSizeF(w, h), QPageSize.Point), QPageLayout.Portrait,
                                     QMarginsF(0, 0, 0, 0)))
        pw.setCreator("DropDial")
        p = QPainter(pw)
        r.render(p, QRectF(0, 0, w, h))
        p.end()
    atomic_save(write, out)


def svg_to(src: str | Path, target: str) -> list[Path]:
    target = target.upper()
    if target == "PDF":
        out = sibling(src, ext=".pdf")
        r = _qt_svg(src)
        if r is not None:
            _svg_to_pdf_qt(r, out)          # stays vector, gradients included
            return [out]
        with _open_svg(src) as doc:
            pdf = fitz.open("pdf", doc.convert_to_pdf())
            atomic_save(lambda dst: pdf.save(str(dst), garbage=3, deflate=True), out)
            pdf.close()
        return [out]
    out = sibling(src, ext="." + target.lower())
    save_image(svg_to_image(src), out, target)
    return [out]


# ---------------------------------------------------------------- e-books (EPUB, MOBI, FB2, XPS, CBZ)
def _open_ebook(src: str | Path):
    try:
        doc = fitz.open(str(src))
    except Exception as e:
        raise UserError("This e-book couldn't be opened (DRM-protected books can't be converted).") from e
    if doc.needs_pass:
        doc.close()
        raise UserError("This e-book is locked.")
    return doc


def ebook_to(src: str | Path, target: str) -> list[Path]:
    target = target.upper()
    with _open_ebook(src) as doc:
        if target == "TXT":
            out = sibling(src, ext=".txt")
            out.write_text("\n\f\n".join(pg.get_text("text") for pg in doc), encoding="utf-8")
            return [out]
        pdf = fitz.open("pdf", doc.convert_to_pdf())
        try:
            toc = doc.get_toc()
            if toc:
                pdf.set_toc(toc)
        except Exception:
            pass
        try:
            if target == "PDF":
                out = sibling(src, ext=".pdf")
                atomic_save(lambda dst: pdf.save(str(dst), garbage=3, deflate=True), out)
            else:
                out = sibling(src, ext=".docx")
                _doc_to_docx(pdf, out)
        finally:
            pdf.close()
    return [out]


# ---------------------------------------------------------------- dispatch
def targets_for(cat: str, src_ext: str) -> list[str]:
    src_ext = src_ext.lower().lstrip(".")
    same = {"jpeg": "JPG", "jpg": "JPG", "tif": "TIFF", "tiff": "TIFF", "heif": "HEIC"}
    cur = same.get(src_ext, src_ext.upper())
    if cat == "image":
        t = [x for x in IMAGE_TARGETS if x != cur]
        if not HEIF_OK:
            t = [x for x in t if x != "HEIC"]
        if not can_trace():
            t = [x for x in t if x != "SVG"]
        return t
    if cat in ("video", "audio"):
        from . import media
        return media.targets_for(cat, src_ext)
    from . import formats
    fam = formats.family(cat)
    if fam is not None:
        return fam.targets(src_ext)
    return {"pdf": PDF_TARGETS, "docx": DOCX_TARGETS, "text": TEXT_TARGETS,
            "vector": VECTOR_TARGETS, "ebook": EBOOK_TARGETS}.get(cat, [])


def convert(src: str | Path, cat: str, target: str) -> list[Path]:
    target = target.upper()
    if cat == "image":
        return image_to(src, target)
    if cat == "pdf":
        if target in ("JPG", "PNG"):
            return pdf_to_images(src, target)
        if target == "TXT":
            return pdf_to_txt(src)
        if target == "DOCX":
            return pdf_to_docx(src)
    if cat == "docx":
        return docx_to_pdf(src) if target == "PDF" else docx_to_txt(src)
    if cat == "text":
        return text_to_pdf(src) if target == "PDF" else text_to_docx(src)
    if cat == "vector":
        return svg_to(src, target)
    if cat == "ebook":
        return ebook_to(src, target)
    if cat == "video":
        from . import media
        return media.convert_video(src, target)
    if cat == "audio":
        from . import media
        return media.convert_audio(src, target)
    from . import formats
    fam = formats.family(cat)
    if fam is not None:
        return fam.convert(Path(src), target)
    raise ValueError(f"Can't convert {cat} to {target}")
