"""Documents and e-mail: Markdown, HTML, saved web pages (MHTML), reStructuredText, .eml.

Every source is first turned into one sanitized HTML tree (no scripts, no forms, no remote
resources: local and embedded images become data: URIs, anything remote is dropped). All
outputs are written from that tree, so a heading is a heading in the PDF, the Word file, the
Markdown and the plain text alike:

* PDF   – PyMuPDF Story (same engine and base CSS as convert.html_to_pdf), with clickable
          links and a bookmark outline built from the headings;
* DOCX  – real Word headings, lists (numbering restarts per list), tables, links, code blocks;
* HTML  – a standalone page with clean CSS and a Content-Security-Policy that forbids any
          network access or script, so opening it never phones home;
* MD / TXT – written from a small block model shared with the DOCX writer.
"""
from __future__ import annotations

import base64
import html
import re
from dataclasses import dataclass, field, replace
from importlib.util import find_spec
from pathlib import Path
from typing import Callable

from . import Family, Tool, register
from ..util import UserError, output_dir, report_progress, sibling

MD_EXTS = {".md", ".markdown"}
HTML_EXTS = {".html", ".htm"}
MHT_EXTS = {".mhtml", ".mht"}
RST_EXTS = {".rst"}
EML_EXTS = {".eml"}

TEXT_LIMIT = 20 * 1024 * 1024        # Markdown / HTML / RST source
MIME_LIMIT = 200 * 1024 * 1024       # MHTML / EML (may carry big attachments)
IMAGE_LIMIT = 30 * 1024 * 1024       # one embedded image
MAX_IMAGES = 500
MAX_PIXELS = 60_000_000
MAX_DEPTH = 120                      # deeper markup is flattened to text

TARGETS = {
    "md": ["PDF", "DOCX", "HTML", "TXT"],
    "markdown": ["PDF", "DOCX", "HTML", "TXT"],
    "html": ["PDF", "DOCX", "MD", "TXT"],
    "htm": ["PDF", "DOCX", "MD", "TXT"],
    "mhtml": ["HTML", "PDF", "DOCX", "MD", "TXT"],
    "mht": ["HTML", "PDF", "DOCX", "MD", "TXT"],
    "rst": ["PDF", "DOCX", "HTML", "MD", "TXT"],
    "eml": ["PDF", "HTML", "DOCX", "TXT"],
}


# ====================================================================== loading
@dataclass
class _Doc:
    title: str
    body: object                 # sanitized lxml element whose children are the content
    page: object = None          # MHTML only: the whole sanitized page (keeps its own CSS)
    lang: str = ""


def _read(src: Path, limit: int) -> bytes:
    try:
        size = src.stat().st_size
    except OSError:
        raise UserError("The file couldn't be opened.") from None
    if size > limit:
        raise UserError(f"This file is too large to convert (over {limit // (1024 * 1024)} MB).")
    return src.read_bytes()


def _decode(raw: bytes, declared: str | None = None) -> str:
    """Bytes -> text: BOM first, then the declared charset, UTF-8, Windows-1252."""
    for bom, enc in ((b"\xef\xbb\xbf", "utf-8"), (b"\xff\xfe", "utf-16-le"), (b"\xfe\xff", "utf-16-be")):
        if raw.startswith(bom):
            return raw[len(bom):].decode(enc, "replace")
    encs = [declared] if declared else []
    for enc in encs + ["utf-8", "cp1252"]:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("latin-1")


def _html_charset(raw: bytes) -> str | None:
    m = re.search(rb"<meta[^>]+charset\s*=\s*[\"']?\s*([A-Za-z0-9_.:-]+)", raw[:8192], re.I)
    return m.group(1).decode("ascii") if m else None


_FRONT_MATTER = re.compile(r"\A---[ \t]*\r?\n(.*?)\r?\n(?:---|\.\.\.)[ \t]*(?:\r?\n|\Z)", re.S)


def _markdown_html(text: str) -> tuple[str, str]:
    """Markdown -> (HTML, front-matter title)."""
    import markdown
    from markdown.extensions import Extension
    from markdown.inlinepatterns import SimpleTagInlineProcessor

    title = ""
    m = _FRONT_MATTER.match(text)
    if m and all(re.match(r"^\s*([\w-]+\s*:|-\s|#|$)", ln) for ln in m.group(1).splitlines()):
        t = re.search(r"^title\s*:\s*(.+)$", m.group(1), re.M)
        title = t.group(1).strip().strip("\"'") if t else ""
        text = text[m.end():]

    class Strike(Extension):            # GitHub's ~~strikethrough~~
        def extendMarkdown(self, md):
            md.inlinePatterns.register(SimpleTagInlineProcessor(r"(~~)(.+?)~~", "del"), "del", 40)

    out = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists", Strike()],
                            output_format="html")
    # GitHub task lists: "- [ ] todo" / "- [x] done"
    out = re.sub(r"<li>(\s*<p>)?\[([ xX])\]\s",
                 lambda m: f"<li>{m.group(1) or ''}{'☑' if m.group(2) in 'xX' else '☐'} ", out)
    return out, title


def _rst_html(text: str) -> str:
    from docutils.core import publish_parts
    settings = {
        "_disable_config": True,          # never read docutils.conf from the current folder
        "file_insertion_enabled": False,  # no .. include::, no raw/csv-table :file:/:url:
        "raw_enabled": False,             # no .. raw:: html
        "report_level": 5, "halt_level": 5, "warning_stream": False,
        "input_encoding": "unicode", "output_encoding": "unicode",
        "embed_stylesheet": False, "syntax_highlight": "none", "math_output": "HTML",
    }
    import warnings
    try:
        with warnings.catch_warnings():          # docutils 0.22 renamed writer_name; both work
            warnings.simplefilter("ignore", PendingDeprecationWarning)
            return publish_parts(text, writer_name="html5", settings_overrides=settings)["html_body"]
    except Exception:
        raise UserError("This reStructuredText file couldn't be read.") from None


def _parse_html(text: str):
    from lxml import etree, html as lh
    text = re.sub(r"\A\s*<\?xml[^>]*\?>", "", text)
    if not text.strip():
        raise UserError("This document is empty.")
    parser = lh.HTMLParser(remove_comments=True, remove_pis=True, no_network=True)
    try:
        root = lh.document_fromstring(text, parser=parser)
    except (etree.ParserError, ValueError):
        raise UserError("This document is empty or isn't valid HTML.") from None
    if any(e.type_name == "ERR_RESOURCE_LIMIT" for e in parser.error_log):
        # libxml2 silently drops whatever lies deeper than 256 levels: refuse, don't lose text
        raise UserError("This document is nested too deeply to convert.")
    return root


def _body(root):
    b = root.find("body")
    return b if b is not None else root


def _title_of(root, fallback: str) -> str:
    t = root.findtext(".//title")
    t = re.sub(r"\s+", " ", t or "").strip()
    if not t:
        h = root.find(".//h1")
        t = re.sub(r"\s+", " ", h.text_content()).strip() if h is not None else ""
    return t or fallback


def _local_images(base: Path) -> Callable[[str], tuple[str, bytes] | None]:
    """Images referenced by a relative path next to the document (never absolute or remote)."""
    from urllib.parse import unquote

    def resolve(src: str):
        s = src.strip()
        if not s or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", s) or s.startswith(("//", "\\\\")):
            return None
        path = unquote(s.split("#")[0].split("?")[0])
        if not path or path.startswith(("/", "\\")):
            return None
        try:
            p = (base / path).resolve()
            if not p.is_file() or p.stat().st_size > IMAGE_LIMIT:
                return None
            return _image_bytes(p.read_bytes(), svg=p.suffix.lower() == ".svg")
        except (OSError, ValueError):
            return None
    return resolve


def _load(src: Path) -> _Doc:
    ext = src.suffix.lower()
    if ext in EML_EXTS:
        return _load_eml(src)
    if ext in MHT_EXTS:
        return _load_mhtml(src)
    raw = _read(src, TEXT_LIMIT)
    title = ""
    if ext in MD_EXTS:
        body_html, title = _markdown_html(_decode(raw))
        text = f"<html><body>{body_html}</body></html>"
    elif ext in RST_EXTS:
        text = f"<html><body>{_rst_html(_decode(raw))}</body></html>"
    else:
        text = _decode(raw, _html_charset(raw))
    root = _parse_html(text)
    lang = root.get("lang", "")
    _sanitize(root, images=_local_images(src.parent))
    return _Doc(title or _title_of(root, src.stem), _body(root), lang=lang)


# ---------------------------------------------------------------- MHTML
def _mime_message(src: Path):
    from email import policy
    from email.parser import BytesParser
    return BytesParser(policy=policy.default).parsebytes(_read(src, MIME_LIMIT))


def _leaves(part):
    """Leaf parts of a MIME tree; an attached e-mail (message/rfc822) counts as one leaf."""
    if part.get_content_maintype() == "multipart":
        for p in part.iter_parts():
            yield from _leaves(p)
    else:
        yield part


def _content_id(part) -> str:
    return (part.get("Content-ID") or "").strip().strip("<>").strip()


def _payload(part) -> bytes:
    try:
        return part.get_payload(decode=True) or b""
    except Exception:
        return b""


def _text_of(part) -> str:
    raw = _payload(part)
    return _decode(raw, part.get_content_charset())


def _load_mhtml(src: Path) -> _Doc:
    from urllib.parse import urljoin
    msg = _mime_message(src)
    parts = list(_leaves(msg))
    start = (msg.get_param("start") or "").strip("<>") if msg.is_multipart() else ""
    root_part = next((p for p in parts if start and _content_id(p) == start), None) or \
        next((p for p in parts if p.get_content_type() == "text/html"), None)
    if root_part is None:
        raise UserError("This saved web page couldn't be read (no web page inside).")
    res: dict[str, object] = {}
    for p in parts:
        if p is root_part:
            continue
        if p.get("Content-Location"):
            res[p["Content-Location"].strip()] = p
        if _content_id(p):
            res["cid:" + _content_id(p)] = p
    base = (root_part.get("Content-Location") or msg.get("Snapshot-Content-Location") or "").strip()

    def find(ref: str, rel_to: str = base):
        ref = html.unescape(ref.strip())
        if ref.lower().startswith("cid:"):
            return res.get("cid:" + ref[4:].strip("<>"))
        return res.get(ref) or (res.get(urljoin(rel_to, ref)) if rel_to else None)

    def image(ref: str):
        p = find(ref)
        return _image_bytes(_payload(p), svg="svg" in p.get_content_type()) if p is not None else None

    def css_url(ref: str, rel_to: str = base) -> str | None:
        p = find(ref, rel_to)
        if p is None:
            return None
        ctype = p.get_content_type()
        if ctype.startswith("image/"):
            r = _image_bytes(_payload(p), svg="svg" in ctype)
            return _data_uri(*r) if r else None
        if ctype.startswith("font/") or ctype in ("application/font-woff", "application/x-font-woff",
                                                    "application/vnd.ms-fontobject", "application/octet-stream"):
            return _data_uri(ctype, _payload(p))
        return None

    def stylesheet(ref: str) -> str | None:
        p = find(ref)
        if p is None or p.get_content_type() != "text/css":
            return None
        loc = (p.get("Content-Location") or base).strip()
        return _clean_css(_text_of(p), lambda u: css_url(u, loc))

    text = _text_of(root_part)
    page = _parse_html(text)
    content = _parse_html(text)
    title = _title_of(content, src.stem)
    _sanitize(page, images=image, page=True, css_url=css_url, stylesheet=stylesheet)
    _sanitize(content, images=image)
    return _Doc(title, _body(content), page=page, lang=content.get("lang", ""))


# ---------------------------------------------------------------- e-mail
def _header(msg, name: str) -> str:
    try:
        v = msg.get(name)
        return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""
    except Exception:            # a malformed header must not stop the conversion
        try:
            from email.header import decode_header, make_header
            return str(make_header(decode_header(msg.get_all(name, [""])[0])))
        except Exception:
            return ""


def _attachment_parts(msg):
    """(part, filename) for every attachment; inline images the body shows are left out."""
    body = msg.get_body(preferencelist=("html", "plain"))
    body_cids = set()
    if body is not None and body.get_content_type() == "text/html":
        body_cids = {html.unescape(c).strip("<>") for c in
                     re.findall(r"cid:([^\s\"'<>)]+)", _text_of(body), re.I)}
    out = []
    for p in _leaves(msg):
        if p is body:
            continue
        disp = p.get_content_disposition()
        try:
            name = p.get_filename() or ""
        except Exception:
            name = ""
        ctype = p.get_content_type()
        if ctype == "message/rfc822":
            if not name:
                try:
                    inner = p.get_content()
                    name = (_header(inner, "Subject") or "message") + ".eml"
                except Exception:
                    name = "message.eml"
        elif disp != "attachment" and not name:
            continue                       # an alternative body part, not a file
        if disp == "inline" and _content_id(p) in body_cids and ctype.startswith("image/"):
            continue
        out.append((p, name))
    return out


_URL = re.compile(r"\bhttps?://[^\s<>\"']+[^\s<>\"'.,;:!?)\]]")


def _plain_to_html(text: str) -> str:
    paras = []
    for para in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip("\n")):
        esc = html.escape(para)
        esc = _URL.sub(lambda m: f'<a href="{m.group(0)}">{m.group(0)}</a>', esc)
        paras.append("<p>" + esc.replace("\n", "<br>") + "</p>")
    return "\n".join(paras)


def _load_eml(src: Path) -> _Doc:
    msg = _mime_message(src)
    if not any(msg.get(h) for h in ("From", "To", "Subject", "Date", "Message-ID")):
        raise UserError("This doesn't look like an e-mail (.eml) file.")
    cids: dict[str, object] = {}
    locs: dict[str, object] = {}
    for p in _leaves(msg):
        if _content_id(p):
            cids[_content_id(p)] = p
        if p.get("Content-Location"):
            locs[p["Content-Location"].strip()] = p

    body = msg.get_body(preferencelist=("html", "plain"))
    body_html = ""
    if body is not None:
        text = _text_of(body)
        body_html = text if body.get_content_type() == "text/html" else _plain_to_html(text)

    def image(ref: str):
        ref = html.unescape(ref.strip())
        p = None
        if ref.lower().startswith("cid:"):
            p = cids.get(ref[4:].strip("<>"))
        elif ref in locs:
            p = locs[ref]
        return _image_bytes(_payload(p), svg="svg" in p.get_content_type()) if p is not None else None

    root = _parse_html(body_html if body_html.strip() else "<html><body></body></html>")
    _sanitize(root, images=image)
    content = _body(root)

    subject = _header(msg, "Subject") or "(no subject)"
    rows = [(k, _header(msg, k)) for k in ("From", "To", "Cc", "Date")]
    names = [_safe_name(n, i, p.get_content_type()) for i, (p, n) in enumerate(_attachment_parts(msg), 1)]
    if names:
        rows.append(("Attachments", ", ".join(names)))
    head = (f"<h1>{html.escape(subject)}</h1><table class=\"email-headers\">"
            + "".join(f"<tr><th>{k}</th><td>{html.escape(v)}</td></tr>" for k, v in rows if v)
            + "</table><hr>")
    from lxml import html as lh
    lead, content.text = content.text, None
    els = lh.fragments_fromstring(head)
    for i, el in enumerate(els):
        content.insert(i, el)
    els[-1].tail = lead                  # body text that came before the first tag
    return _Doc(subject, content)


# ====================================================================== sanitizing
DROP = {"script", "noscript", "template", "iframe", "frame", "frameset", "object", "embed",
        "applet", "param", "input", "button", "select", "textarea", "option", "optgroup",
        "datalist", "output", "link", "meta", "base", "canvas", "audio", "video", "source",
        "track", "map", "area", "dialog", "portal", "slot"}
CONTENT_DROP = {"style", "svg", "head", "title", "math"}
CONTENT_ATTRS = {"href", "src", "alt", "title", "colspan", "rowspan", "start", "value", "class",
                 "lang", "dir", "data-orig"}
URL_DROP_ATTRS = {"srcset", "action", "formaction", "background", "poster", "ping", "longdesc",
                  "cite", "codebase", "data", "manifest", "lowsrc", "dynsrc", "imagesizes"}
SAFE_SCHEMES = {"http", "https", "mailto", "tel"}


def _scheme(url: str) -> str | None:
    s = re.sub(r"[\x00-\x20]", "", html.unescape(url))
    m = re.match(r"^([a-zA-Z][a-zA-Z0-9+.-]*):", s)
    return m.group(1).lower() if m else None


def _clean_css(css: str, resolve: Callable[[str], str | None] | None = None) -> str:
    """No @import and no url() outside the page: archived resources become data: URIs."""
    css = re.sub(r"@import\b[^;]*;?", "", css, flags=re.I)
    css = re.sub(r"expression\s*\(", "(", css, flags=re.I)

    def url(m):
        ref = m.group(2).strip()
        if ref.lower().startswith("data:"):
            return m.group(0) if not re.match(r"data:\s*text/(html|javascript)", ref, re.I) else 'url("data:,")'
        new = resolve(ref) if resolve else None
        return f'url("{new}")' if new else 'url("data:,")'
    return re.sub(r"url\(\s*(['\"]?)(.*?)\1\s*\)", url, css, flags=re.I | re.S)


def _is_layout_table(t) -> bool:
    if (t.get("role") or "").lower() in ("presentation", "none"):
        return True
    if any(True for _ in t.iterdescendants("table")):
        return True
    rows = [r for r in t.iter("tr")]
    return not rows or max(sum(1 for c in r if c.tag in ("td", "th")) for r in rows) <= 1


def _sanitize(root, *, images=None, page: bool = False, css_url=None, stylesheet=None) -> None:
    """Make a parsed page safe in place. Content mode (default) also strips presentation."""
    drop = DROP if page else DROP | CONTENT_DROP
    for el in [e for e in root.iter() if isinstance(e.tag, str) and e.tag.lower() in drop]:
        if page and el.tag == "link" and stylesheet and "stylesheet" in (el.get("rel") or "").lower():
            css = stylesheet(el.get("href") or "")
            if css is not None:
                from lxml import html as lh
                st = lh.Element("style")
                st.text = css
                el.addprevious(st)
        el.drop_tree()
    if not page:
        for t in [t for t in root.iter("table") if _is_layout_table(t)]:
            for el in list(t.iter("thead", "tbody", "tfoot", "colgroup", "col", "caption")):
                if el.tag in ("colgroup", "col"):
                    el.drop_tree()
                else:
                    el.drop_tag()
            for el in t.iter("tr", "td", "th"):
                el.tag = "div"
            t.tag = "div"
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        tag = el.tag.lower()
        for name in list(el.attrib):
            low = name.lower()
            local = low.rsplit(":", 1)[-1]
            if low.startswith("on") or local in URL_DROP_ATTRS or (not page and low not in CONTENT_ATTRS):
                del el.attrib[name]
            elif local == "href":
                v = el.get(name)
                sch = _scheme(v)
                if tag == "a":
                    if sch is not None and sch not in SAFE_SCHEMES:
                        del el.attrib[name]
                elif not v.strip().startswith("#"):        # svg <use>/<image>: same page only
                    del el.attrib[name]
            elif local == "src" and tag != "img":
                del el.attrib[name]
            elif low == "style":
                el.set(name, _clean_css(el.get(name), css_url))
        if page and tag == "style":
            el.text = _clean_css(el.text or "", css_url)
    if not page:                  # a link that lost its unsafe target is plain text
        for a in [a for a in root.iter("a") if not a.get("href")]:
            a.drop_tag()
    count = 0
    cache: dict[str, str | None] = {}
    for img in list(root.iter("img")):
        src = (img.get("src") or "").strip()
        if src not in cache:
            res = None
            if count < MAX_IMAGES and src:
                if src.lower().startswith("data:"):
                    d = _parse_data_uri(src)
                    res = _image_bytes(d[1], svg="svg" in d[0]) if d else None
                elif images is not None:
                    res = images(src)
            cache[src] = _data_uri(*res) if res else None
            count += res is not None
        if cache[src]:
            if not src.lower().startswith("data:") and not page:
                img.set("data-orig", src)
            img.set("src", cache[src])
        else:                                 # remote or missing: show what a browser shows
            img.tail = (img.get("alt") or "") + (img.tail or "")
            img.drop_tree()


# ---------------------------------------------------------------- images
def _parse_data_uri(s: str) -> tuple[str, bytes] | None:
    from urllib.parse import unquote_to_bytes
    m = re.match(r"data:([^,;]*)((?:;[^,;]*)*),(.*)\Z", s.strip(), re.S | re.I)
    if not m or len(m.group(3)) > IMAGE_LIMIT * 4 // 3 + 16:
        return None
    mime = (m.group(1) or "text/plain").strip().lower()
    try:
        if ";base64" in m.group(2).lower():
            payload = re.sub(r"\s", "", m.group(3))
            return mime, base64.b64decode(payload + "=" * (-len(payload) % 4))
        return mime, unquote_to_bytes(m.group(3))
    except (ValueError, TypeError):
        return None


def _data_uri(mime: str, data: bytes) -> str:
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"


def _image_bytes(raw: bytes, svg: bool = False) -> tuple[str, bytes] | None:
    """Validate an image and return (mime, bytes) that Word and MuPDF can both show.
    PNG/JPEG/GIF are kept byte for byte; other formats (and SVG) become PNG."""
    import io
    if not raw or len(raw) > IMAGE_LIMIT:
        return None
    head = raw[:2048].lstrip().lower()
    if svg or head.startswith(b"<svg") or (head.startswith(b"<?xml") and b"<svg" in head):
        return _svg_png(raw)
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(raw))
        w, h = im.size
        if w * h > MAX_PIXELS or w < 1 or h < 1:
            return None
        if im.format in ("PNG", "JPEG", "GIF"):
            return {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif"}[im.format], raw
        from .. import convert
        out = io.BytesIO()
        convert.open_image(io.BytesIO(raw)).save(out, "PNG")
        return "image/png", out.getvalue()
    except Exception:
        return None


def _svg_png(raw: bytes) -> tuple[str, bytes] | None:
    if b"<!entity" in raw[:4096].lower():           # no entity tricks in untrusted SVG
        return None
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        import fitz  # type: ignore
    try:
        with fitz.open(stream=raw, filetype="svg") as doc:
            page = doc[0]
            side = max(page.rect.width, page.rect.height, 1)
            zoom = max(0.1, min(2.0, 2000 / side))
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=True)
            return "image/png", pix.tobytes("png")
    except Exception:
        return None


# ====================================================================== block model
@dataclass
class Run:
    text: str = ""
    b: bool = False
    i: bool = False
    u: bool = False
    s: bool = False
    code: bool = False
    sup: bool = False
    sub: bool = False
    href: str | None = None
    kind: str = "text"                 # text | br | img
    img: bytes | None = None
    src: str = ""                      # original image reference (kept in Markdown)
    alt: str = ""

    def fmt(self):
        return (self.b, self.i, self.u, self.s, self.code, self.sup, self.sub, self.href)


@dataclass
class Cell:
    runs: list[Run]
    header: bool = False
    colspan: int = 1


@dataclass
class Block:
    kind: str                          # p | h | li | pre | table | hr
    runs: list[Run] = field(default_factory=list)
    level: int = 0                     # heading level, or list depth (0 = not in a list)
    ordered: bool = False
    number: int = 1
    list_id: int = 0
    quote: int = 0
    cont: bool = False                 # paragraph that continues a list item
    text: str = ""                     # pre
    lang: str = ""
    rows: list[list[Cell]] = field(default_factory=list)
    kv: bool = False                   # e-mail header table: "From: …" lines


BLOCK_TAGS = {"p", "div", "section", "article", "main", "header", "footer", "nav", "aside",
              "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "menu", "dir", "li", "pre",
              "blockquote", "table", "thead", "tbody", "tfoot", "tr", "td", "th", "caption",
              "hr", "dl", "dt", "dd", "figure", "figcaption", "address", "details", "summary",
              "center", "body", "html", "form", "fieldset", "legend", "hgroup", "search"}
INLINE_FMT = {"b": "b", "strong": "b", "i": "i", "em": "i", "cite": "i", "dfn": "i", "var": "i",
              "u": "u", "ins": "u", "s": "s", "strike": "s", "del": "s", "code": "code",
              "kbd": "code", "samp": "code", "tt": "code", "sup": "sup", "sub": "sub"}
_WS = re.compile(r"[ \t\n\r\f\v]+")


@dataclass(frozen=True)
class _Ctx:
    lists: tuple = ()                  # (ordered | None for a plain indent, list id)
    quote: int = 0


def _tidy(runs: list[Run]) -> list[Run]:
    """Merge runs, collapse spaces across them, trim each line; [] if nothing visible."""
    out: list[Run] = []
    for r in runs:
        if r.kind == "text":
            if not r.text:
                continue
            if out and out[-1].kind == "text" and out[-1].fmt() == r.fmt():
                out[-1].text += r.text
                continue
        out.append(r)
    while out and out[0].kind == "br":
        out.pop(0)
    while out and out[-1].kind == "br":
        out.pop()
    at_start = True
    for k, r in enumerate(out):
        if r.kind == "text":
            if at_start:
                r.text = r.text.lstrip(" ")
            if r.text:
                at_start = r.text.endswith(" ")
        elif r.kind == "br":
            for prev in reversed(out[:k]):
                if prev.kind != "text":
                    break
                prev.text = prev.text.rstrip(" ")
                if prev.text:
                    break
            at_start = True
        else:
            at_start = False
    for r in reversed(out):
        if r.kind != "text":
            break
        r.text = r.text.rstrip(" ")
        if r.text:
            break
    out = [r for r in out if r.kind != "text" or r.text]
    if not any(r.kind == "img" or (r.kind == "text" and r.text.strip()) for r in out):
        return []
    return out


def _int(v, default: int, lo: int = -10**6, hi: int = 10**6) -> int:
    try:
        return max(lo, min(hi, int(str(v).strip())))
    except (TypeError, ValueError):
        return default


class _Walker:
    """Sanitized HTML tree -> flat list of Blocks."""

    def __init__(self):
        self.out: list[Block] = []
        self.nlists = 0

    def run(self, root) -> list[Block]:
        self._container(root, _Ctx(), 0)
        return self.out

    def _para(self, runs, ctx: _Ctx):
        runs = _tidy(runs)
        if runs:
            self.out.append(Block("p", runs, level=len(ctx.lists), cont=bool(ctx.lists), quote=ctx.quote))

    def _container(self, el, ctx: _Ctx, depth: int):
        buf: list[Run] = []
        if el.text:
            buf += _text(el.text, {})
        for ch in el:
            if isinstance(ch.tag, str):
                if depth > MAX_DEPTH:
                    buf += _text(ch.text_content(), {})
                elif ch.tag in BLOCK_TAGS:
                    self._para(buf, ctx)
                    buf = []
                    self._block(ch, ctx, depth + 1)
                else:
                    self._inline(ch, {}, buf, depth + 1)
            if ch.tail:
                buf += _text(ch.tail, {})
        self._para(buf, ctx)

    def _block(self, el, ctx: _Ctx, depth: int):
        tag = el.tag
        lvl = len(ctx.lists)
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            runs = _tidy(self._inlines(el, depth))
            if runs:
                self.out.append(Block("h", runs, level=int(tag[1]), quote=ctx.quote))
        elif tag == "pre":
            for br in el.iter("br"):
                br.tail = "\n" + (br.tail or "")
            text = el.text_content()
            text = text[1:] if text.startswith("\n") else text
            m = re.search(r"\b(?:language|lang)-([\w+#.-]+)", " ".join(
                (e.get("class") or "") for e in el.iter() if isinstance(e.tag, str)))
            self.out.append(Block("pre", text=text.rstrip("\n"), lang=m.group(1) if m else "",
                                  level=lvl, cont=bool(lvl), quote=ctx.quote))
        elif tag in ("ul", "ol", "menu", "dir"):
            self._list(el, ctx, depth, tag == "ol")
        elif tag == "li":
            self.nlists += 1
            self._item(el, replace(ctx, lists=ctx.lists + ((False, self.nlists),)), 1, depth)
        elif tag == "blockquote":
            self._container(el, replace(ctx, quote=ctx.quote + 1), depth)
        elif tag == "table":
            self._table(el, ctx, depth)
        elif tag == "hr":
            self.out.append(Block("hr", quote=ctx.quote))
        elif tag == "dt":
            self._para(self._inlines(el, depth, {"b": True}), ctx)
        elif tag == "dd":
            self._container(el, replace(ctx, lists=ctx.lists + ((None, 0),)), depth)
        else:
            self._container(el, ctx, depth)

    def _list(self, el, ctx: _Ctx, depth: int, ordered: bool):
        self.nlists += 1
        sub = replace(ctx, lists=ctx.lists + ((ordered, self.nlists),))
        n = _int(el.get("start"), 1) if ordered else 1
        for ch in el:
            if not isinstance(ch.tag, str):
                continue
            if ch.tag == "li":
                n = _int(ch.get("value"), n) if ordered else n
                self._item(ch, sub, n, depth + 1)
                n += 1
            else:                                   # stray content, e.g. a list in a list
                self._block(ch, sub, depth + 1) if ch.tag in BLOCK_TAGS else \
                    self._para(self._inlines(ch, depth + 1), sub)

    def _item(self, li, ctx: _Ctx, number: int, depth: int):
        start = len(self.out)
        self._container(li, ctx, depth)
        lvl = len(ctx.lists)
        ordered, lid = ctx.lists[-1]
        first = self.out[start] if len(self.out) > start else None
        if first is not None and first.kind == "p" and first.cont and first.level == lvl:
            first.kind, first.cont = "li", False
        else:
            first = Block("li", level=lvl, quote=ctx.quote)
            self.out.insert(start, first)
        first.ordered, first.number, first.list_id = bool(ordered), number, lid

    def _table(self, el, ctx: _Ctx, depth: int):
        cap = el.find("caption")
        if cap is not None:
            self._para(self._inlines(cap, depth, {"b": True}), ctx)
        rows = []
        for tr, head in _table_rows(el):
            cells = [Cell(self._cell(c, depth), header=head or c.tag == "th",
                          colspan=_int(c.get("colspan"), 1, 1, 63))
                     for c in tr if isinstance(c.tag, str) and c.tag in ("td", "th")]
            if cells:
                rows.append(cells)
        if rows:
            self.out.append(Block("table", rows=rows, level=len(ctx.lists), quote=ctx.quote,
                                  kv="email-headers" in (el.get("class") or "")))

    def _cell(self, c, depth: int) -> list[Run]:
        w = _Walker()
        w._container(c, _Ctx(), depth + 1)
        runs: list[Run] = []
        for k, b in enumerate(w.out):
            if k:
                runs.append(Run(kind="br"))
            if b.kind == "li":
                runs.append(Run(text=f"{b.number}. " if b.ordered else "• "))
            if b.kind == "pre":
                for j, line in enumerate(b.text.split("\n")):
                    runs += ([Run(kind="br")] if j else []) + [Run(text=line, code=True)]
            elif b.kind == "table":
                for j, row in enumerate(b.rows):
                    runs += [Run(kind="br")] if j else []
                    for cell in row:
                        runs += cell.runs + [Run(text=" ")]
            else:
                runs += b.runs
        return runs

    def _inlines(self, el, depth: int, fmt: dict | None = None) -> list[Run]:
        buf: list[Run] = []
        fmt = fmt or {}
        if el.text:
            buf += _text(el.text, fmt)
        for ch in el:
            self._inline(ch, fmt, buf, depth + 1)
            if ch.tail:
                buf += _text(ch.tail, fmt)
        return buf

    def _inline(self, el, fmt: dict, buf: list[Run], depth: int):
        tag = el.tag
        if not isinstance(tag, str):
            return
        if tag == "br":
            buf.append(Run(kind="br"))
            return
        if tag == "img":
            d = _parse_data_uri(el.get("src") or "")
            if d:
                buf.append(Run(kind="img", img=d[1], src=el.get("data-orig") or el.get("src"),
                               alt=el.get("alt") or "", href=fmt.get("href")))
            return
        f = dict(fmt)
        if tag in INLINE_FMT:
            f[INLINE_FMT[tag]] = True
        elif tag == "a" and el.get("href"):
            f["href"] = el.get("href").strip()
        elif tag == "span" and "literal" in (el.get("class") or "").split():
            f["code"] = True                       # docutils inline literal
        if depth > MAX_DEPTH:
            buf += _text(el.text_content(), f)
            return
        block = tag in BLOCK_TAGS
        if block:
            buf.append(Run(kind="br"))
        if tag == "q":
            buf.append(Run(text="“", **_fmtargs(f)))
        if el.text:
            buf += _text(el.text, f)
        for ch in el:
            self._inline(ch, f, buf, depth + 1)
            if ch.tail:
                buf += _text(ch.tail, f)
        if tag == "q":
            buf.append(Run(text="”", **_fmtargs(f)))
        if block:
            buf.append(Run(kind="br"))


def _fmtargs(f: dict) -> dict:
    return {k: v for k, v in f.items() if k in ("b", "i", "u", "s", "code", "sup", "sub", "href")}


def _text(s: str, fmt: dict) -> list[Run]:
    s = _WS.sub(" ", s)
    return [Run(text=s, **_fmtargs(fmt))] if s else []


def _table_rows(t):
    for ch in t:
        if not isinstance(ch.tag, str):
            continue
        if ch.tag == "tr":
            yield ch, False
        elif ch.tag in ("thead", "tbody", "tfoot"):
            for tr in ch:
                if isinstance(tr.tag, str) and tr.tag == "tr":
                    yield tr, ch.tag == "thead"


def _blocks(doc: _Doc) -> list[Block]:
    blocks = _Walker().run(doc.body)
    if not blocks:
        raise UserError("This document has no content to convert.")
    return blocks


def _link_groups(runs: list[Run]):
    """Split runs into (href, [runs]) groups of consecutive runs sharing one link."""
    groups: list[tuple[str | None, list[Run]]] = []
    for r in runs:
        h = r.href if r.kind != "br" else None
        if groups and groups[-1][0] == h and h is not None:
            groups[-1][1].append(r)
        else:
            groups.append((h, [r]))
    return groups


def _new_list(b: Block, top: int) -> bool:
    """A top-level item that starts a different list than the one before it."""
    return b.kind == "li" and b.level == 1 and b.list_id != top


def _is_web(href: str | None) -> bool:
    return bool(href) and _scheme(href) in SAFE_SCHEMES


# ====================================================================== writers
def _atomic(write, out: Path) -> None:
    from .. import convert
    convert.atomic_save(write, out)


# ---------------------------------------------------------------- TXT
def _write_txt(doc: _Doc, out: Path) -> None:
    def inline(runs: list[Run], indent: str) -> str:
        parts = []
        for href, group in _link_groups(runs):
            text = "".join(r.text if r.kind == "text" else f"\n{indent}" if r.kind == "br"
                           else (f"[{r.alt}]" if r.alt else "") for r in group)
            target = re.sub(r"^mailto:", "", href or "")
            if _is_web(href) and target.rstrip("/") != text.strip().rstrip("/"):
                text += f" ({target})"
            parts.append(text)
        return "".join(parts)

    lines: list[str] = []
    prev, top = None, 0
    for b in _blocks(doc):
        indent = "    " * b.quote + "    " * (b.level - (1 if b.kind == "li" else 0) if b.level else 0)
        if b.kind == "li":
            indent = "    " * b.quote + "  " * (b.level - 1)
            mark = f"{b.number}. " if b.ordered else "• "
            text = indent + mark + inline(b.runs, indent + " " * len(mark))
        elif b.kind == "h":
            text = inline(b.runs, "")
        elif b.kind == "pre":
            text = "\n".join(indent + ln for ln in b.text.split("\n"))
        elif b.kind == "table":
            rows = []
            for row in b.rows:
                cells = [inline(c.runs, "").replace("\n", " ") for c in row]
                rows.append(indent + (f"{cells[0]}: {' '.join(cells[1:])}" if b.kv and len(cells) > 1
                                      else "\t".join(cells)))
            text = "\n".join(rows)
        elif b.kind == "hr":
            text = indent + "* * *"
        else:
            text = indent + inline(b.runs, indent)
        listy = prev is not None and prev.level and b.level and b.kind in ("li", "p", "pre") \
            and prev.kind in ("li", "p", "pre") and not _new_list(b, top)
        if b.kind == "li" and b.level == 1:
            top = b.list_id
        if prev is not None:
            lines.append("" if not listy else None)
        lines.append(text)
        prev = b
    body = "\n".join(ln for ln in lines if ln is not None).rstrip() + "\n"
    _atomic(lambda dst: Path(dst).write_text(body, encoding="utf-8"), out)


# ---------------------------------------------------------------- Markdown
def _md_escape(text: str) -> str:
    text = re.sub(r"([\\`*\[\]<>|])", r"\\\1", text)
    return re.sub(r"(?<![A-Za-z0-9])_|_(?![A-Za-z0-9])", r"\\_", text)


def _md_inline(runs: list[Run], cont: str = "  \n") -> str:
    def one(r: Run) -> str:
        if r.kind == "br":
            return cont
        if r.kind == "img":
            return f"![{_md_escape(r.alt)}]({_md_url(r.src)})"
        if r.code:
            ticks = "`" * (max((len(m) for m in re.findall(r"`+", r.text)), default=0) + 1)
            pad = " " if r.text.startswith("`") or r.text.endswith("`") else ""
            core = f"{ticks}{pad}{r.text}{pad}{ticks}"
        else:
            core = _md_escape(r.text)
        lead = core[: len(core) - len(core.lstrip(" "))]
        trail = core[len(core.rstrip(" ")):]
        core = core.strip(" ")
        if not core:
            return lead + trail
        if r.sup:
            core = f"<sup>{core}</sup>"
        if r.sub:
            core = f"<sub>{core}</sub>"
        if r.s:
            core = f"~~{core}~~"
        if r.i:
            core = f"*{core}*"
        if r.b:
            core = f"**{core}**"
        return lead + core + trail

    out = []
    for href, group in _link_groups(runs):
        inner = "".join(one(r) for r in group)
        if href:
            lead = inner[: len(inner) - len(inner.lstrip(" "))]
            trail = inner[len(inner.rstrip(" ")):]
            inner = f"{lead}[{inner.strip(' ')}]({_md_url(href)}){trail}"
        out.append(inner)
    return "".join(out)


def _md_url(url: str) -> str:
    url = url.strip()
    return f"<{url}>" if re.search(r"[\s()<>]", url) and not url.startswith("data:") else url


def _write_md(doc: _Doc, out: Path) -> None:
    lines: list[str | None] = []
    prev, top = None, 0
    for b in _blocks(doc):
        q = "> " * b.quote
        ind = "    " * (b.level - 1 if b.kind == "li" else b.level)
        if b.kind == "h":
            text = q + "#" * min(b.level, 6) + " " + _md_inline(b.runs, " ")
        elif b.kind == "li":
            mark = f"{b.number}. " if b.ordered else "- "
            text = q + ind + mark + _md_inline(b.runs, "  \n" + q + ind + "    ")
        elif b.kind == "pre":
            fence = "`" * max(3, max((len(m) for m in re.findall(r"`{3,}", b.text)), default=0) + 1)
            body = [fence + b.lang] + b.text.split("\n") + [fence]
            text = "\n".join(q + ind + ln if ln else q.rstrip() for ln in body)
        elif b.kind == "table":
            ncols = max(sum(c.colspan for c in row) for row in b.rows)
            rows = []
            for row in b.rows:
                cells = []
                for c in row:
                    cells.append(_md_inline(c.runs, "<br>").replace("\n", " "))
                    cells += [""] * (c.colspan - 1)
                rows.append(cells + [""] * (ncols - len(cells)))
            fmt = [q + ind + "| " + " | ".join(r) + " |" for r in rows]
            text = "\n".join([fmt[0], q + ind + "|" + "|".join(["---"] * ncols) + "|"] + fmt[1:])
        elif b.kind == "hr":
            text = q + "---"
        else:
            body = _md_inline(b.runs, "  \n" + q + ind)
            body = re.sub(r"^(#{1,6}\s|[-+*]\s|\d+[.)]\s|=+\s*$|-+\s*$)", r"\\\1", body)
            text = q + ind + body
        tight = prev is not None and b.kind == "li" and prev.kind == "li" and \
            (prev.list_id == b.list_id or prev.level != b.level) and not _new_list(b, top)
        if b.kind == "li" and b.level == 1:
            top = b.list_id
        if prev is not None:
            lines.append(None if tight else q.rstrip() if prev.quote and b.quote else "")
        lines.append(text)
        prev = b
    body = "\n".join(ln for ln in lines if ln is not None).rstrip() + "\n"
    _atomic(lambda dst: Path(dst).write_text(body, encoding="utf-8"), out)


# ---------------------------------------------------------------- HTML
HTML_CSS = """
:root { color-scheme: light dark; --fg: #1d1d1f; --muted: #5f6368; --bg: #fff; --line: #d0d7de;
        --code: #f3f4f6; --link: #0b57d0; }
@media (prefers-color-scheme: dark) {
  :root { --fg: #e6e6e6; --muted: #a0a4a8; --bg: #16171a; --line: #3a3d42; --code: #23252a;
          --link: #8ab4f8; } }
html { background: var(--bg); color: var(--fg); }
body { margin: 0; padding: 2.5rem 1rem 4rem; font: 16px/1.6 system-ui, -apple-system, "Segoe UI",
       Roboto, "Noto Sans", "Helvetica Neue", Arial, sans-serif; }
main { max-width: 46rem; margin: 0 auto; overflow-wrap: break-word; }
h1, h2, h3, h4, h5, h6 { line-height: 1.25; margin: 1.6em 0 .6em; }
h1 { font-size: 2em; margin-top: 0; } h2 { font-size: 1.5em; border-bottom: 1px solid var(--line);
     padding-bottom: .25em; } h3 { font-size: 1.25em; }
a { color: var(--link); }
img { max-width: 100%; height: auto; }
code, kbd, samp, pre { font-family: ui-monospace, SFMono-Regular, Menlo, Consolas, "Liberation Mono",
                       monospace; font-size: .9em; }
code { background: var(--code); padding: .1em .35em; border-radius: 4px; }
pre { background: var(--code); padding: .9em 1em; border-radius: 6px; overflow-x: auto; line-height: 1.45; }
pre code { background: none; padding: 0; }
blockquote { margin: 1em 0; padding: 0 1em; color: var(--muted); border-left: 4px solid var(--line); }
table { border-collapse: collapse; margin: 1em 0; display: block; overflow-x: auto; }
th, td { border: 1px solid var(--line); padding: .4em .75em; text-align: left; vertical-align: top; }
th { background: var(--code); }
hr { border: 0; border-top: 1px solid var(--line); margin: 2em 0; }
table.email-headers { display: table; border: 0; margin: 0 0 1em; }
table.email-headers th, table.email-headers td { border: 0; background: none; padding: .15em .8em .15em 0; }
table.email-headers th { color: var(--muted); font-weight: 600; }
"""
CSP = ("default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src data:; "
       "base-uri 'none'; form-action 'none'")


def _inner_html(el) -> str:
    from lxml import html as lh
    return html.escape(el.text or "") + "".join(lh.tostring(c, encoding="unicode", method="html")
                                                for c in el)


def _write_html(doc: _Doc, out: Path) -> None:
    from lxml import html as lh
    if doc.page is not None:                     # a saved web page keeps its own look
        page = doc.page
        head = page.find("head")
        if head is None:
            head = lh.Element("head")
            page.insert(0, head)
        for i, (k, v) in enumerate((("charset", "utf-8"), ("http-equiv", "Content-Security-Policy"))):
            m = lh.Element("meta")
            m.set(k, v)
            if k == "http-equiv":
                m.set("content", CSP)
            head.insert(i, m)
        text = lh.tostring(page, encoding="unicode", method="html", doctype="<!DOCTYPE html>")
    else:
        for el in doc.body.iter("img"):
            el.attrib.pop("data-orig", None)
        lang = f' lang="{html.escape(doc.lang)}"' if doc.lang else ""
        text = (f"<!DOCTYPE html>\n<html{lang}>\n<head>\n<meta charset=\"utf-8\">\n"
                f"<meta http-equiv=\"Content-Security-Policy\" content=\"{CSP}\">\n"
                "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
                f"<title>{html.escape(doc.title)}</title>\n<style>{HTML_CSS}</style>\n</head>\n"
                f"<body>\n<main>\n{_inner_html(doc.body).strip()}\n</main>\n</body>\n</html>\n")
    _atomic(lambda dst: Path(dst).write_text(text, encoding="utf-8"), out)


# ---------------------------------------------------------------- PDF
PDF_CSS = """
body { font-size: 10.5pt; line-height: 1.4; }
h1 { font-size: 20pt; margin: 0 0 8pt 0; } h2 { font-size: 15pt; margin: 14pt 0 6pt 0; }
h3 { font-size: 12.5pt; margin: 12pt 0 4pt 0; } h4, h5, h6 { font-size: 11pt; margin: 10pt 0 4pt 0; }
p { margin: 0 0 7pt 0; }
ul, ol { margin: 0 0 7pt 0; } li { margin: 0 0 2pt 0; } li p { margin: 0 0 2pt 0; }
code { font-family: monospace; font-size: 9pt; background-color: #f0f0f0; }
pre { font-family: monospace; font-size: 9pt; white-space: pre-wrap; background-color: #f4f4f4;
      padding: 6pt; margin: 0 0 8pt 0; }
pre code { background-color: #f4f4f4; }
blockquote { margin: 0 0 8pt 10pt; padding: 0 0 0 8pt; border-left: 2pt solid #bbbbbb; color: #444444; }
table { margin: 4pt 0 9pt 0; }
th { border: 1px solid #888; padding: 3pt 5pt; font-size: 10pt; background-color: #eeeeee; font-weight: bold; }
a { color: #0b57d0; }
hr { border-top: 1px solid #cccccc; margin: 8pt 0; }
table.email-headers td, table.email-headers th { border: 0; background-color: #ffffff; padding: 1pt 8pt 1pt 0;
      text-align: left; }
table.email-headers th { color: #555555; }
"""


def _write_pdf(doc: _Doc, out: Path) -> None:
    from .. import convert
    fitz = convert.fitz
    _blocks(doc)                             # refuses an empty document with a clear message
    for n, h in enumerate(doc.body.iter("h1", "h2", "h3", "h4", "h5", "h6")):
        h.set("id", f"h{n}")
    for el in doc.body.iter("img"):
        el.attrib.pop("data-orig", None)
    story = fitz.Story(html=_inner_html(doc.body), user_css=convert.CSS + PDF_CSS)
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    heads: list[tuple[int, str, int]] = []

    def rect(n, filled):
        report_progress(min(0.9, 0.3 + n * 0.02))
        return mediabox, where, None

    def position(p):
        if p.heading and p.open_close & 1 and p.text:
            heads.append((p.heading, re.sub(r"\s+", " ", p.text).strip(), p.page_num))

    pdf = story.write_with_links(rect, position)
    try:
        toc, last = [], 0
        for lvl, text, page in heads:
            lvl = min(lvl, last + 1)
            toc.append([lvl, text, page])
            last = lvl
        try:
            if toc:
                pdf.set_toc(toc)
        except Exception:
            pass
        pdf.set_metadata({"title": doc.title, "creator": "DropDial", "producer": "DropDial"})
        try:
            pdf.subset_fonts()         # CJK/Indic fallback fonts are megabytes when embedded whole
        except Exception:
            pass
        _atomic(lambda dst: pdf.save(str(dst), garbage=3, deflate=True), out)
    finally:
        pdf.close()


# ---------------------------------------------------------------- DOCX
_PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
                  "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                  "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


class _Docx:
    def __init__(self, title: str):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt, RGBColor
        self.d = Document()
        self.d.core_properties.title = title[:250]
        st = self.d.styles
        code = st.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = st["Normal"]
        code.font.name = "Consolas"
        code.font.size = Pt(9.5)
        pf = code.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(2), Pt(8), 1.0
        self._shade(code.element.get_or_add_pPr(), "F3F4F6")
        ic = st.add_style("Inline Code", WD_STYLE_TYPE.CHARACTER)
        ic.font.name = "Consolas"
        ic.font.size = Pt(9.5)
        ic.font.color.rgb = RGBColor(0x8B, 0x1A, 0x3B)
        link = st.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        link.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        link.font.underline = True
        self.nums: dict[int, int] = {}

    @staticmethod
    def _shade(ppr, fill: str):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        ppr.insert_element_before(shd, *_PPR_AFTER_BDR[1:])

    def style(self, name: str, level: int) -> str:
        return name if level <= 1 else f"{name} {min(level, 3)}"

    def restart(self, p, style: str, list_id: int, start: int):
        """Each ordered list gets its own numbering instance, so it starts at its own number."""
        if list_id not in self.nums:
            numbering = self.d.part.numbering_part.element
            style_num = self.d.styles[style].element.pPr.numPr.numId.val
            abstract = numbering.num_having_numId(style_num).abstractNumId.val
            num = numbering.add_num(abstract)
            num.add_lvlOverride(ilvl=0).add_startOverride(start)
            self.nums[list_id] = num.numId
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = self.nums[list_id]

    def runs(self, p, runs: list[Run], max_width: float = 6.0, bold: bool = False):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from urllib.parse import quote
        for href, group in _link_groups(runs):
            hl = None
            if _is_web(href):
                url = quote(href, safe=":/?#[]@!$&'()*+,;=%~-._")
                hl = OxmlElement("w:hyperlink")
                hl.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
                p._p.append(hl)
            for r in group:
                run = self.run(p, r, max_width, bold)
                if run is not None and hl is not None:
                    if r.kind == "text" and not r.code:
                        run.style = "Hyperlink"
                    hl.append(run._r)

    def run(self, p, r: Run, max_width: float, bold: bool):
        import io
        from docx.shared import Inches
        if r.kind == "br":
            run = p.add_run()
            run.add_break()
            return run
        if r.kind == "img":
            from PIL import Image
            try:
                w_px = Image.open(io.BytesIO(r.img)).size[0]
                run = p.add_run()
                run.add_picture(io.BytesIO(r.img), width=Inches(max(0.1, min(max_width, w_px / 96))))
                return run
            except Exception:
                return p.add_run(r.alt) if r.alt else None
        run = p.add_run(r.text)
        run.bold = (r.b or bold) or None
        run.italic = r.i or None
        run.underline = r.u or None
        if r.s:
            run.font.strike = True
        if r.sup:
            run.font.superscript = True
        if r.sub:
            run.font.subscript = True
        if r.code:
            run.style = "Inline Code"
        return run

    def table(self, b: Block):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        ncols = min(63, max(sum(c.colspan for c in row) for row in b.rows))
        t = self.d.add_table(rows=len(b.rows), cols=ncols)
        if not b.kv:
            t.style = "Table Grid"
        width = max(0.6, 6.0 / ncols)
        for ri, row in enumerate(b.rows):
            ci = 0
            for cell in row:
                if ci >= ncols:
                    break
                span = min(cell.colspan, ncols - ci)
                dc = t.cell(ri, ci)
                if span > 1:
                    dc = dc.merge(t.cell(ri, ci + span - 1))
                self.runs(dc.paragraphs[0], cell.runs, width * span, bold=cell.header or (b.kv and ci == 0))
                ci += span
            if row and all(c.header for c in row) and ri == 0:
                trpr = t.rows[0]._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                h.set(qn("w:val"), "true")
                trpr.append(h)
        self.d.add_paragraph()           # keeps two tables in a row from merging into one

    def write(self, blocks: list[Block], out: Path):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches
        d = self.d
        for b in blocks:
            if b.kind == "h":
                self.runs(d.add_paragraph(style=f"Heading {min(b.level, 6)}"), b.runs)
            elif b.kind == "li":
                style = self.style("List Number" if b.ordered else "List Bullet", b.level)
                p = d.add_paragraph(style=style)
                if b.ordered:
                    self.restart(p, style, b.list_id, b.number)
                if b.level > 3:
                    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * b.level)
                self.runs(p, b.runs)
            elif b.kind == "pre":
                p = d.add_paragraph(style="Code Block")
                p.add_run(b.text)          # python-docx turns \n and \t into breaks and tabs
                if b.level or b.quote:
                    p.paragraph_format.left_indent = Inches(0.25 * (b.level + b.quote) + 0.25)
            elif b.kind == "table":
                self.table(b)
            elif b.kind == "hr":
                p = d.add_paragraph()
                bdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "BBBBBB")):
                    bottom.set(qn(f"w:{k}"), v)
                bdr.append(bottom)
                p._p.get_or_add_pPr().insert_element_before(bdr, *_PPR_AFTER_BDR)
            else:
                if b.quote:
                    style = "Quote"
                elif b.cont:
                    style = self.style("List Continue", b.level)
                else:
                    style = None
                p = d.add_paragraph(style=style)
                if b.cont and b.level > 3:
                    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * b.level)
                self.runs(p, b.runs)
        _atomic(lambda dst: d.save(str(dst)), out)


def _write_docx(doc: _Doc, out: Path) -> None:
    _Docx(doc.title).write(_blocks(doc), out)


# ====================================================================== entry points
WRITERS = {"PDF": _write_pdf, "DOCX": _write_docx, "HTML": _write_html, "MD": _write_md, "TXT": _write_txt}


def targets(src_ext: str) -> list[str]:
    return list(TARGETS.get(src_ext.lower().lstrip("."), []))


def convert(src: Path, target: str) -> list[Path]:
    src = Path(src)
    target = target.upper()
    if target not in targets(src.suffix):
        raise UserError(f"Can't turn a {src.suffix or 'file'} into {target}.")
    report_progress(0.05)
    doc = _load(src)
    report_progress(0.3)
    out = sibling(src, ext="." + target.lower())
    try:
        WRITERS[target](doc, out)
    except UserError:
        raise
    except RecursionError:
        raise UserError("This document is nested too deeply to convert.") from None
    report_progress(1.0)
    return [out]


# ---------------------------------------------------------------- attachments
_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}


def _safe_name(name: str, index: int, ctype: str = "") -> str:
    """A file name that is safe on Windows, macOS and Linux and can't leave its folder."""
    import mimetypes
    import unicodedata
    name = unicodedata.normalize("NFC", name or "")
    name = name.replace("\\", "/").split("/")[-1]                  # drop any directory part
    name = re.sub(r'[\x00-\x1f\x7f<>:"|?*]', "_", name)
    name = name.strip().strip(".").strip()
    stem, dot, ext = name.rpartition(".")
    if not dot:
        stem, ext = name, ""
    if stem.split(".")[0].lower() in _RESERVED:
        name = "_" + name
    if not name or set(name) <= {"_", " ", "."}:
        guess = mimetypes.guess_extension(ctype or "") or ""
        name = f"attachment {index}{'.eml' if ctype == 'message/rfc822' else guess}"
    if len(name.encode("utf-8")) > 180:                             # most file systems: 255 bytes
        ext = ("." + ext) if dot and len(ext) <= 16 else ""
        keep = name[: len(name) - len(ext)]
        while len((keep + ext).encode("utf-8")) > 180:
            keep = keep[:-1]
        name = keep.rstrip() + ext
    return name


def save_attachments(src: Path) -> list[Path]:
    src = Path(src)
    if src.suffix.lower() not in EML_EXTS:
        raise UserError("Only e-mails (.eml) have attachments to save.")
    msg = _mime_message(src)
    if not any(msg.get(h) for h in ("From", "To", "Subject", "Date", "Message-ID")):
        raise UserError("This doesn't look like an e-mail (.eml) file.")
    parts = _attachment_parts(msg)
    if not parts:
        raise UserError("This e-mail has no attachments.")
    base = output_dir(src)
    folder = base / f"{src.stem} attachments"
    k = 1
    while folder.exists():
        folder = base / f"{src.stem} attachments ({k})"
        k += 1
    folder.mkdir(parents=True)
    root = folder.resolve()
    taken: set[str] = set()
    outs = []
    for i, (part, name) in enumerate(parts, 1):
        ctype = part.get_content_type()
        if ctype == "message/rfc822":
            try:
                data = part.get_content().as_bytes()
            except Exception:
                data = b"".join(m.as_bytes() for m in part.iter_parts())
        else:
            data = _payload(part)
        name = _safe_name(name, i, ctype)
        stem, dot, ext = name.rpartition(".")
        stem, ext = (stem, "." + ext) if dot and stem else (name, "")
        cand, n = name, 1
        while cand.lower() in taken:
            cand = f"{stem} ({n}){ext}"
            n += 1
        taken.add(cand.lower())
        path = folder / cand
        if path.resolve().parent != root:          # belt and braces against traversal
            raise UserError("An attachment has an unsafe name.")
        with open(path, "xb") as f:                # never overwrite anything
            f.write(data)
        outs.append(path)
        report_progress(i / len(parts))
    return outs


# ---------------------------------------------------------------- registration
def _have(mod: str) -> bool:
    try:
        return find_spec(mod) is not None
    except (ImportError, ValueError):
        return False


_EXTS = set(HTML_EXTS | MHT_EXTS)
if _have("markdown"):
    _EXTS |= MD_EXTS
if _have("docutils"):
    _EXTS |= RST_EXTS

HINTS = {"PDF": "Typeset for printing", "DOCX": "Editable in Word", "HTML": "Opens in any browser",
         "MD": "Clean Markdown text", "TXT": "Plain text, no markup"}

register(Family(
    category="markup",
    label="Document",
    exts=_EXTS,
    targets=targets,
    convert=convert,
    hints=dict(HINTS),
    available=lambda: _have("lxml"),
))

register(Family(
    category="email",
    label="E-mail",
    exts=set(EML_EXTS),
    targets=targets,
    convert=convert,
    hints={**HINTS, "PDF": "Headers + message", "Attachments": "Save every attachment"},
    tools=[Tool("Attachments", "⤓", "Save every attachment", lambda paths: save_attachments(paths[0]),
                working="Saving attachments…")],
    available=lambda: _have("lxml"),
))
