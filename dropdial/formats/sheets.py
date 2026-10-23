"""Spreadsheets and plain tables: XLSX/XLSM, XLS (read only), ODS, CSV and TSV.

Everything goes through one in-memory shape, `read_tables(path) -> [(sheet name, rows)]`, where a
row is a list of plain Python values (str, int, float, bool, date, datetime, time, timedelta or
None). CSV/TSV give strings; `infer_types` turns "42", "3.50" and "2024-05-01" into numbers and
dates (leading zeros such as "00123" stay text). Writers then turn tables into each target.
"""
from __future__ import annotations

import csv
import datetime as dt
import html
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter
from importlib.util import find_spec
from pathlib import Path

from ..util import UserError, report_progress, sibling
from . import Family, Tool, register

Table = tuple[str, list[list]]

MAX_CELLS = 5_000_000          # read limit per file
MAX_PDF_CELLS = 400_000        # ~ 5,000 pages
MAX_DOCX_CELLS = 150_000       # Word itself crawls beyond this
MAX_UNZIPPED = 1 << 30         # 1 GB of XML inside one workbook
MAX_TEXT_BYTES = 1 << 30
XLSX_MAX_TEXT = 32_767         # Excel's limit per cell

TEXT_EXTS = {".csv", ".tsv"}
XLS_OK = find_spec("xlrd") is not None
EXTS = {".xlsx", ".xlsm", ".ods", ".csv", ".tsv"} | ({".xls"} if XLS_OK else set())

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _openpyxl_ok() -> bool:
    return find_spec("openpyxl") is not None


# ---------------------------------------------------------------- values
def _num(v: float) -> int | float:
    """Excel keeps 15 significant digits: 0.1+0.2 -> 0.3, 5.0 -> 5."""
    if not math.isfinite(v):
        return v
    v = float(f"{v:.15g}")
    return int(v) if v.is_integer() and abs(v) < 1e15 else v


def _duration(td: dt.timedelta) -> str:
    secs = td.total_seconds()
    sign, secs = ("-" if secs < 0 else ""), abs(secs)
    h, rem = divmod(int(round(secs)), 3600)
    return f"{sign}{h}:{rem // 60:02d}:{rem % 60:02d}"


def _iso(v, sep="T") -> str:
    if isinstance(v, dt.datetime):
        if v.tzinfo is None and v.time() == dt.time(0):
            return v.date().isoformat()
        return v.isoformat(sep=sep, timespec="microseconds" if v.microsecond else "seconds")
    return v.isoformat()


def cell_text(v, sep: str = " ") -> str:
    """How a value reads in a text format (CSV, Markdown, PDF…). Dates are ISO."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        n = _num(v)
        return str(n) if isinstance(n, int) else f"{n:.15g}"
    if isinstance(v, (dt.date, dt.time)):
        return _iso(v, sep)
    if isinstance(v, dt.timedelta):
        return _duration(v)
    return str(v)


def _json_value(v):
    if v is None or isinstance(v, (bool, str, int)):
        return v
    if isinstance(v, float):
        return _num(v) if math.isfinite(v) else None
    if isinstance(v, (dt.date, dt.time)):
        return _iso(v)
    return cell_text(v)


_INT = re.compile(r"-?[1-9]\d{0,14}|0")
_DEC = re.compile(r"-?(?:0|[1-9]\d*)\.(\d+)")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?")


def _infer(v):
    if not isinstance(v, str):
        return v
    if _INT.fullmatch(v):
        return int(v)
    if _DEC.fullmatch(v) and len(v.lstrip("-").replace(".", "").lstrip("0")) <= 15:
        return float(v)
    try:
        if _DATE.fullmatch(v):
            return dt.date.fromisoformat(v)
        if _DATETIME.fullmatch(v):
            return dt.datetime.fromisoformat(v)
    except ValueError:
        pass
    return v


def infer_types(rows: list[list]) -> list[list]:
    """Text cells that are clearly numbers or ISO dates become numbers and dates; the first row
    (the header) and anything with leading zeros, spaces or more than 15 digits stays text."""
    return rows[:1] + [[_infer(v) for v in row] for row in rows[1:]]


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _decimals(s: str) -> int:
    m = _DEC.fullmatch(s)
    return len(m.group(1)) if m else 0


def _clean(s: str) -> str:
    return _ILLEGAL.sub("", s)


# ---------------------------------------------------------------- reading
def _finish(rows: list[list]) -> list[list]:
    """Drop trailing empty cells/rows and pad to a rectangle."""
    for row in rows:
        while row and (row[-1] is None or row[-1] == ""):
            row.pop()
        for i, v in enumerate(row):
            if v == "":
                row[i] = None
    while rows and not rows[-1]:
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    for row in rows:
        row.extend([None] * (width - len(row)))
    return rows


def _too_big():
    raise UserError("This table is too big to convert (over 5 million cells).")


def _check_zip(path: Path) -> zipfile.ZipFile:
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        with open(path, "rb") as f:
            head = f.read(8)
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            raise UserError("This workbook is password-protected. Remove the password in Excel first.") from None
        raise UserError("This spreadsheet is damaged or isn't really a spreadsheet.") from None
    total = 0
    for info in zf.infolist():
        total += info.file_size
        if total > MAX_UNZIPPED or (info.file_size > 100 << 20 and info.file_size > 400 * max(info.compress_size, 1)):
            zf.close()
            raise UserError("This spreadsheet is too big to convert.")
    return zf


def _read_xlsx(path: Path) -> list[Table]:
    _check_zip(path).close()
    import warnings
    from openpyxl import load_workbook
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb_v = load_workbook(path, read_only=True, data_only=True)
            wb_f = load_workbook(path, read_only=True, data_only=False)
    except UserError:
        raise
    except Exception as e:
        raise UserError("This spreadsheet couldn't be opened (it may be damaged).") from e
    try:
        tables, cells = [], 0
        for ws_v, ws_f in zip(wb_v.worksheets, wb_f.worksheets):
            ws_v.reset_dimensions()   # a lying <dimension> (A1:XFD1048576) must not pad every row
            ws_f.reset_dimensions()
            rows = []
            for rv, rf in zip(ws_v.iter_rows(values_only=True), ws_f.iter_rows(values_only=True)):
                row = list(rv)
                for i, v in enumerate(row):
                    if v is None and i < len(rf) and rf[i] is not None:
                        f = rf[i]
                        row[i] = getattr(f, "text", f)     # formula with no saved result: its text
                cells += len(row)
                if cells > MAX_CELLS:
                    _too_big()
                rows.append(row)
            tables.append((ws_v.title, _finish(rows), ws_v.sheet_state == "visible"))
    except UserError:
        raise
    except Exception as e:
        raise UserError("This spreadsheet couldn't be read (it may be damaged).") from e
    finally:
        wb_v.close()
        wb_f.close()
    return _visible(tables)


def _visible(tables: list) -> list[Table]:
    """Hidden sheets are left out, unless every sheet is hidden."""
    shown = [(n, r) for n, r, vis in tables if vis]
    return shown or [(n, r) for n, r, _ in tables]


def _read_xls(path: Path) -> list[Table]:
    import xlrd
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as e:
        raise UserError("This .xls file couldn't be opened (it may be damaged or password-protected).") from e
    tables, cells = [], 0
    try:
        for i in range(book.nsheets):
            sh = book.sheet_by_index(i)
            cells += sh.nrows * sh.ncols
            if cells > MAX_CELLS:
                _too_big()
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    t, v = sh.cell_type(r, c), sh.cell_value(r, c)
                    if t == xlrd.XL_CELL_NUMBER:
                        v = _num(v)
                    elif t == xlrd.XL_CELL_DATE:
                        try:
                            d = xlrd.xldate_as_datetime(v, book.datemode)
                            v = d.time() if v < 1 else d
                        except Exception:
                            v = _num(v)
                    elif t == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    elif t == xlrd.XL_CELL_ERROR:
                        v = xlrd.error_text_from_code.get(v, "#ERROR")
                    elif t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        v = None
                    row.append(v)
                rows.append(row)
            tables.append((sh.name, _finish(rows), sh.visibility == 0))
            book.unload_sheet(i)
    except UserError:
        raise
    except Exception as e:
        raise UserError("This .xls file couldn't be read (it may be damaged).") from e
    finally:
        book.release_resources()
    return _visible(tables)


# ODS (OpenDocument): content.xml read with defusedxml, streaming row by row
_NS_T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_NS_O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_NS_X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_PT = re.compile(r"-?P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:(\d+(?:\.\d+)?)S)?")


def _ods_text(cell) -> str | None:
    def walk(e, out):
        if e.text:
            out.append(e.text)
        for ch in e:
            tag = ch.tag
            if tag == _NS_X + "s":
                out.append(" " * int(ch.get(_NS_X + "c", "1") or 1))
            elif tag == _NS_X + "tab":
                out.append("\t")
            elif tag == _NS_X + "line-break":
                out.append("\n")
            elif tag not in (_NS_O + "annotation", _NS_X + "note"):
                walk(ch, out)
            if ch.tail:
                out.append(ch.tail)
    paras = []
    for p in cell:
        if p.tag in (_NS_X + "p", _NS_X + "h"):
            out: list[str] = []
            walk(p, out)
            paras.append("".join(out))
    return "\n".join(paras) or None


def _ods_value(cell):
    vt = cell.get(_NS_O + "value-type")
    try:
        if vt in ("float", "percentage", "currency"):
            return _num(float(cell.get(_NS_O + "value")))
        if vt == "boolean":
            return cell.get(_NS_O + "boolean-value") == "true"
        if vt == "date":
            s = cell.get(_NS_O + "date-value")
            return dt.datetime.fromisoformat(s) if "T" in s else dt.date.fromisoformat(s)
        if vt == "time":
            m = _PT.fullmatch(cell.get(_NS_O + "time-value"))
            d, h, mi, s = (float(x or 0) for x in m.groups())
            td = dt.timedelta(days=d, hours=h, minutes=mi, seconds=s)
            if td < dt.timedelta(days=1):
                return (dt.datetime.min + td).time()
            return td
    except (TypeError, ValueError, AttributeError):
        pass
    return _ods_text(cell)


def _read_ods(path: Path) -> list[Table]:
    from defusedxml import ElementTree as DET
    zf = _check_zip(path)
    try:
        if b"opendocument.spreadsheet" not in zf.read("mimetype")[:100]:
            raise UserError("This file isn't an OpenDocument spreadsheet.")
        tables: list = []
        rows: list[list] = []
        name, cells, pend_rows = "", 0, 0
        with zf.open("content.xml") as f:
            for event, el in DET.iterparse(f, events=("start", "end"), forbid_dtd=True):
                tag = el.tag
                if event == "start":
                    if tag == _NS_T + "table":
                        name, rows, pend_rows = el.get(_NS_T + "name") or f"Sheet{len(tables) + 1}", [], 0
                    continue
                if tag == _NS_T + "table-row":
                    row, pend = [], 0
                    for c in el:
                        if c.tag not in (_NS_T + "table-cell", _NS_T + "covered-table-cell"):
                            continue
                        rep = int(c.get(_NS_T + "number-columns-repeated", "1") or 1)
                        v = _ods_value(c)
                        if v is None:
                            pend += rep          # trailing empties are often repeated 16384×
                            continue
                        if len(row) + pend + rep > 16384:
                            raise UserError("This spreadsheet has too many columns.")
                        row += [None] * pend + [v] * rep
                        pend = 0
                    rep = int(el.get(_NS_T + "number-rows-repeated", "1") or 1)
                    if not row:
                        pend_rows += rep
                    else:
                        cells += len(row) * (rep + pend_rows)
                        if cells > MAX_CELLS:
                            _too_big()
                        rows += [[] for _ in range(pend_rows)] + [list(row) for _ in range(rep)]
                        pend_rows = 0
                    el.clear()
                elif tag == _NS_T + "table":
                    tables.append((name, _finish(rows), True))
                    el.clear()
    except UserError:
        raise
    except Exception as e:
        raise UserError("This spreadsheet couldn't be read (it may be damaged).") from e
    finally:
        zf.close()
    return _visible(tables)


def _decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in text:
        raise UserError("This doesn't look like a text table.")
    return text


def _sniff(sample: str, prefer_tab: bool) -> str:
    """Pick the delimiter that splits the sample into the most consistent column count."""
    lines = sample.splitlines()
    if len(lines) > 1:
        sample = "\n".join(lines[:-1])   # the last line may be cut off
    order = ["\t", ",", ";", "|"] if prefer_tab else [",", ";", "\t", "|"]
    best, best_score = order[0], 0.0
    for d in order:
        try:
            counts = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d) if r][:500]
        except csv.Error:
            continue
        if not counts:
            continue
        mode, n = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        score = round(n / len(counts), 2)      # ties go to the earlier, more common delimiter
        if score > best_score:
            best, best_score = d, score
    return best


def _read_csv(path: Path) -> list[Table]:
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise UserError("This file is too big to convert.")
    text = _decode(path.read_bytes())
    delim = _sniff(text[:65536], path.suffix.lower() == ".tsv")
    csv.field_size_limit(max(csv.field_size_limit(), 16 << 20))
    rows, cells = [], 0
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delim):
            cells += len(row)
            if cells > MAX_CELLS:
                _too_big()
            rows.append(row)
    except csv.Error as e:
        raise UserError("This file couldn't be read as a table.") from e
    return [(path.stem, _finish(rows))]


_OLE = b"\xd0\xcf\x11\xe0"


def _is_text(path: Path, head: bytes | None = None) -> bool:
    """CSV/TSV, or an ".xls" that is really tab-separated text (many web apps save those)."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        return True
    if ext != ".xls":
        return False
    if head is None:
        with open(path, "rb") as f:
            head = f.read(8)
    return not head.startswith((_OLE, b"PK")) and not head.lstrip().startswith((b"<", b"\xef\xbb\xbf<"))


def read_tables(path: str | Path) -> list[Table]:
    """Every sheet of a spreadsheet or CSV/TSV as (name, rows); rows are rectangular lists of
    values (trailing empty rows/columns trimmed). CSV/TSV cells stay strings: see infer_types."""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        with open(path, "rb") as f:
            head = f.read(8)
    except OSError as e:
        raise UserError("This file couldn't be opened.") from e
    if _is_text(path, head):
        return _read_csv(path)
    if head.startswith(b"PK"):
        with _check_zip(path) as zf:
            is_ods = "mimetype" in zf.namelist() and b"opendocument" in zf.read("mimetype")[:100]
        return _read_ods(path) if is_ods else _read_xlsx(path)
    if head.startswith(_OLE):
        if ext == ".xls":
            return _read_xls(path)
        raise UserError("This workbook is password-protected. Remove the password in Excel first.")
    if ext == ".xls" and head.lstrip().lower().startswith((b"<", b"\xef\xbb\xbf<")):
        raise UserError("This .xls file is really a web page. Open it in Excel and save it as XLSX.")
    raise UserError("This spreadsheet is damaged or isn't really a spreadsheet.")


# ---------------------------------------------------------------- shaping
def _empty(rows) -> bool:
    return not any(v is not None and v != "" for row in rows for v in row)


def _tidy(rows: list[list]) -> list[list]:
    """Drop leading empty rows and columns (the table starts where the data starts)."""
    start = next((i for i, r in enumerate(rows) if any(v is not None for v in r)), len(rows))
    rows = rows[start:]
    lead = min((next((i for i, v in enumerate(r) if v is not None), len(r)) for r in rows), default=0)
    return [r[lead:] for r in rows] if lead else rows


def _tables(src: Path, typed: bool) -> list[Table]:
    tables = [(n, r) for n, r in read_tables(src) if not _empty(r)]
    if not tables:
        raise UserError("This spreadsheet is empty.")
    if typed and _is_text(src):
        tables = [(n, infer_types(r)) for n, r in tables]
    return tables


def _headers(row: list) -> list[str]:
    out, seen = [], Counter()
    for i, v in enumerate(row):
        h = cell_text(v).strip() or f"Column {i + 1}"
        seen[h] += 1
        out.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    return out


def _safe_name(name: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", name).strip(" .")
    return s[:60] or "Sheet"


def _sheet_names(names: list[str]) -> list[str]:
    """Valid, unique Excel sheet names (≤ 31 chars, no []:*?/\\)."""
    out, used = [], set()
    for n in names:
        base = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", n).strip("' ")[:31] or "Sheet"
        cand, i = base, 2
        while cand.lower() in used:
            suf = f" ({i})"
            cand, i = base[:31 - len(suf)] + suf, i + 1
        used.add(cand.lower())
        out.append(cand)
    return out


def _per_sheet_outputs(src: Path, tables: list[Table], ext: str) -> list[tuple[Path, Table]]:
    if len(tables) == 1:
        return [(sibling(src, ext=ext), tables[0])]
    return [(sibling(src, f"- {_safe_name(n)}", ext), (n, r)) for n, r in tables]


def _display_width(s: str) -> int:
    line = max(s.split("\n"), key=len) if s else ""
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line)


# ---------------------------------------------------------------- writers
def _save(write, out: Path) -> None:
    from ..convert import atomic_save
    atomic_save(write, out)


def write_xlsx(tables: list[Table], out: str | Path, infer: bool | list[bool] = False) -> Path:
    """Write tables to an .xlsx file, one sheet each, with a bold frozen header row and sized
    columns. Strings are always stored as text, never as formulas. With `infer` (per table or for
    all), text such as "3.50" or "2024-05-01" becomes a number or date, keeping its look (0.00)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    out = Path(out)
    flags = infer if isinstance(infer, list) else [infer] * len(tables)
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    names = _sheet_names([n for n, _ in tables])
    for k, (name, (_, rows), guess) in enumerate(zip(names, tables, flags)):
        ws = wb.create_sheet(title=name)
        widths: dict[int, int] = {}
        for row in rows[:2000]:
            for i, v in enumerate(row):
                widths[i] = max(widths.get(i, 0), _display_width(cell_text(v)))
        for i, w in widths.items():
            ws.column_dimensions[get_column_letter(i + 1)].width = min(max(w + 2, 8), 60)
        if len(rows) > 1:
            ws.freeze_panes = "A2"
        for r, row in enumerate(rows):
            line = []
            for v in row:
                if v is None:
                    line.append(None)
                    continue
                fmt = None
                if guess and r > 0 and isinstance(v, str):
                    typed = _infer(v)
                    if isinstance(typed, float) and 0 < _decimals(v) <= 10:
                        fmt = "0." + "0" * _decimals(v)          # "3.50" shows as 3.50
                    v = typed
                if isinstance(v, float) and not math.isfinite(v):
                    v = cell_text(v)
                elif isinstance(v, dt.datetime) and v.tzinfo is not None:
                    v = v.replace(tzinfo=None)
                elif not isinstance(v, (str, int, float, dt.date, dt.time, dt.timedelta)):
                    v = cell_text(v)
                if isinstance(v, str):
                    v = _clean(v)[:XLSX_MAX_TEXT]
                c = WriteOnlyCell(ws, value=v)
                if isinstance(v, str):
                    c.data_type = "s"                        # "=cmd|…" stays text, never a formula
                    if _NUMERIC_TEXT.fullmatch(v):
                        fmt = "@"                            # "00123" stays text when edited
                if fmt:
                    c.number_format = fmt
                if r == 0:
                    c.font = bold
                line.append(c)
            ws.append(line)
        report_progress((k + 1) / len(tables))
    if not tables:
        wb.create_sheet("Sheet1")
    _save(wb.save, out)
    return out


_NUMERIC_TEXT = re.compile(r"[+-]?\d+(?:[.,]\d+)?")


def _write_delimited(rows: list[list], out: Path, delim: str, bom: bool) -> None:
    def write(dst):
        with open(dst, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as f:
            w = csv.writer(f, delimiter=delim, lineterminator="\r\n")
            for row in rows:
                w.writerow([cell_text(v) for v in row])
    _save(write, out)


def _write_json(rows: list[list], out: Path) -> None:
    rows = _tidy(rows)
    keys = _headers(rows[0]) if rows else []
    data = [dict(zip(keys, (_json_value(v) for v in row))) for row in rows[1:]
            if any(v is not None for v in row)]
    _save(lambda dst: Path(dst).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                                           encoding="utf-8"), out)


def _numeric_cols(rows: list[list]) -> list[bool]:
    """Columns whose body is all numbers (right-aligned in tables)."""
    if len(rows) < 2:
        return [False] * (len(rows[0]) if rows else 0)
    cols = []
    for i in range(len(rows[0])):
        vals = [_infer(r[i]) for r in rows[1:] if r[i] is not None]
        cols.append(bool(vals) and all(_is_number(v) for v in vals))
    return cols


def _md_cell(v) -> str:
    s = cell_text(v).replace("\\", "\\\\").replace("|", "\\|")
    return s.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")


def _write_md(tables: list[Table], out: Path) -> None:
    parts = []
    for name, rows in tables:
        rows = _tidy(rows)
        if len(tables) > 1:
            parts.append(f"## {_md_cell(name)}\n")
        num = _numeric_cols(rows)
        lines = ["| " + " | ".join(_md_cell(h) for h in _headers(rows[0])) + " |",
                 "| " + " | ".join("---:" if n else "---" for n in num) + " |"]
        lines += ["| " + " | ".join(_md_cell(v) for v in row) + " |" for row in rows[1:]]
        parts.append("\n".join(lines) + "\n")
    _save(lambda dst: Path(dst).write_text("\n".join(parts), encoding="utf-8"), out)


def _h(v) -> str:
    return html.escape(cell_text(v)).replace("\n", "<br>")


HTML_CSS = """
:root { color-scheme: light dark; --bg: #fff; --fg: #1d2330; --line: #d5dae3; --head: #eef1f6; --zebra: #f8f9fb; }
@media (prefers-color-scheme: dark) { :root { --bg: #16181d; --fg: #e6e8ee; --line: #343944; --head: #22262e; --zebra: #1b1e24; } }
body { margin: 24px; background: var(--bg); color: var(--fg); font: 14px/1.45 system-ui, -apple-system, "Segoe UI", sans-serif; }
h1 { font-size: 20px; margin: 0 0 16px; } h2 { font-size: 16px; margin: 28px 0 10px; }
.wrap { overflow-x: auto; border: 1px solid var(--line); border-radius: 8px; }
table { border-collapse: collapse; min-width: 100%; }
th, td { padding: 6px 10px; border-bottom: 1px solid var(--line); text-align: left; vertical-align: top; white-space: pre-wrap; }
th { position: sticky; top: 0; background: var(--head); font-weight: 600; }
tbody tr:nth-child(even) { background: var(--zebra); }
td.n, th.n { text-align: right; font-variant-numeric: tabular-nums; }
"""


def _table_html(rows: list[list], row_ids: int | None = None,
                num: list[bool] | None = None) -> str:
    num = num if num is not None else _numeric_cols(rows)
    cls = [' class="n"' if n else "" for n in num]
    out = ["<table><thead><tr>"]
    out += [f"<th{cls[i]}>{_h(v)}</th>" for i, v in enumerate(rows[0])]
    out.append("</tr></thead><tbody>")
    for k, row in enumerate(rows[1:]):
        rid = f' id="r{row_ids + k}"' if row_ids is not None else ""
        out.append(f"<tr{rid}>" + "".join(f"<td{cls[i]}>{_h(v)}</td>" for i, v in enumerate(row)) + "</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def _write_html(tables: list[Table], title: str, out: Path) -> None:
    body = [f"<h1>{html.escape(title)}</h1>"]
    for name, rows in tables:
        rows = _tidy(rows)
        if len(tables) > 1:
            body.append(f"<h2>{html.escape(name)}</h2>")
        body.append(f'<div class="wrap">{_table_html(rows)}</div>')
    doc = (f'<!doctype html>\n<html><head><meta charset="utf-8">'
           f'<meta name="viewport" content="width=device-width, initial-scale=1">'
           f"<title>{html.escape(title)}</title><style>{HTML_CSS}</style></head>\n<body>\n"
           + "\n".join(body) + "\n</body></html>\n")
    _save(lambda dst: Path(dst).write_text(doc, encoding="utf-8"), out)


# PDF: MuPDF's Story lays out HTML (with font fallback for Hindi, Japanese, emoji…). We page it
# ourselves so every page repeats the header row; wide sheets go landscape and are scaled down.
PDF_CSS = """
* { font-family: sans-serif; }
h2 { font-size: 13pt; margin: 0 0 6pt 0; }
table { border-collapse: collapse; }
th, td { border: 0.6pt solid #9aa3b2; padding: 2pt 4pt; font-size: 8.5pt; text-align: left; vertical-align: top; }
th { background-color: #e9edf3; font-weight: bold; }
td.n, th.n { text-align: right; }
"""
PDF_FONT = 8.5
PDF_MARGIN = 36
PDF_MAX_CELL = 1500     # characters shown per cell
PDF_MIN_SCALE = 0.5


def _fitz():
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover
        import fitz  # type: ignore
    return fitz


def _natural_width(rows: list[list]) -> float:
    cols: dict[int, float] = {}
    for row in rows[:3000]:
        for i, v in enumerate(row):
            w = min(_display_width(cell_text(v)[:200]) * 0.56 * PDF_FONT, 260)
            cols[i] = max(cols.get(i, 0), w)
    return sum(w + 10 for w in cols.values()) or 100


def _write_pdf(tables: list[Table], out: Path) -> None:
    fitz = _fitz()
    total = sum(len(r) for _, r in tables)

    buf = io.BytesIO()

    def layout():
        writer = fitz.DocumentWriter(buf)
        done = 0
        try:
            for name, rows in tables:
                rows = _tidy(rows)
                rows = [[v[:PDF_MAX_CELL] if isinstance(v, str) else v for v in r] for r in rows]
                num = _numeric_cols(rows)
                head, body = rows[0], rows[1:]
                natural = _natural_width(rows)
                portrait = fitz.paper_rect("a4")
                page = portrait if natural <= portrait.width - 2 * PDF_MARGIN else fitz.paper_rect("a4-l")
                aw, ah = page.width - 2 * PDF_MARGIN, page.height - 2 * PDF_MARGIN
                scale = min(1.0, max(PDF_MIN_SCALE, aw / natural))
                where = fitz.Rect(0, 0, aw / scale, ah / scale)
                mat = fitz.Matrix(scale, 0, 0, scale, PDF_MARGIN, PDF_MARGIN)
                title = f"<h2>{html.escape(name)}</h2>" if len(tables) > 1 else ""

                def story(i, k, with_title):
                    h = (title if with_title else "") + _table_html([head] + body[i:i + k], row_ids=i, num=num)
                    return fitz.Story(html=h, user_css=PDF_CSS)

                def place(i, k, with_title):
                    """(rows that fit completely, the laid-out story if all k fit)."""
                    s = story(i, k, with_title)
                    more, _ = s.place(where)
                    if not more:
                        return min(k, len(body) - i), s
                    closed = set()
                    s.element_positions(lambda p: closed.add(p.id) if p.id and p.open_close & 2 else None)
                    return sum(1 for j in range(i, i + k) if f"r{j}" in closed), None

                i, k, first = 0, 40, True
                while first or i < len(body):
                    while True:
                        n, s = place(i, k, first)
                        if s is None or i + k >= len(body):
                            break
                        k *= 2
                    if n == 0 and body:
                        # one row taller than a page: let it run over as many pages as it needs
                        s = story(i, 1, first)
                        more = True
                        while more:
                            dev = writer.begin_page(page)
                            more, _ = s.place(where)
                            s.draw(dev, mat)
                            writer.end_page()
                        n = 1
                    else:
                        if s is None:
                            s = story(i, n, first)
                            s.place(where)
                        dev = writer.begin_page(page)
                        s.draw(dev, mat)
                        writer.end_page()
                    i += n
                    k = n + 1            # pages hold about as many rows as the last one
                    first = False
                    done += n
                    report_progress(done / max(total, 1))
        finally:
            writer.close()

    layout()
    # every page carries its own copy of the fonts: one cleanup pass shares them (≈10× smaller)
    pdf = fitz.open("pdf", buf.getvalue())
    try:
        try:
            pdf.subset_fonts()        # the CJK fallback font alone is ~1.5 MB
        except Exception:  # pragma: no cover
            pass
        _save(lambda dst: pdf.save(str(dst), garbage=3, deflate=True), out)
    finally:
        pdf.close()


def _docx_text(s: str) -> str:
    """Run content: text with line breaks and tabs as Word elements."""
    parts = re.split(r"(\r\n|\n|\r|\t)", _clean(s))
    out = []
    for t in parts:
        if t == "\t":
            out.append("<w:tab/>")
        elif t in ("\n", "\r", "\r\n"):
            out.append("<w:br/>")
        elif t:
            out.append(f'<w:t xml:space="preserve">{html.escape(t, quote=False)}</w:t>')
    return "".join(out)


def _write_docx(tables: list[Table], out: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches

    tables = [(n, _tidy(r)) for n, r in tables]
    doc = Document()
    sec = doc.sections[0]
    widest = max(len(r[0]) for _, r in tables)
    if widest > 6:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    half_pts = 20 if widest <= 6 else 17 if widest <= 12 else 14
    total, done = sum(len(r) for _, r in tables), 0
    for name, rows in tables:
        if len(tables) > 1:
            doc.add_heading(name, level=2)
        num = _numeric_cols(rows)
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        # rows are built as XML in one go: python-docx's cell API is quadratic on big tables
        xml = []
        for r, row in enumerate(rows):
            xml.append("<w:tr>" + ("<w:trPr><w:tblHeader/></w:trPr>" if r == 0 else ""))
            rpr = f'<w:rPr>{"<w:b/>" if r == 0 else ""}<w:sz w:val="{half_pts}"/></w:rPr>'
            for c, v in enumerate(row):
                ppr = '<w:pPr><w:jc w:val="right"/></w:pPr>' if num[c] else ""
                text = _docx_text(cell_text(v))
                run = f"<w:r>{rpr}{text}</w:r>" if text else ""
                xml.append(f'<w:tc><w:tcPr><w:tcW w:w="0" w:type="auto"/></w:tcPr><w:p>{ppr}{run}</w:p></w:tc>')
            xml.append("</w:tr>")
            done += 1
            if done % 500 == 0:
                report_progress(done / total)
        frag = parse_xml(f"<w:tbl {nsdecls('w')}>{''.join(xml)}</w:tbl>")
        for tr in list(frag):
            table._tbl.append(tr)
    _save(lambda dst: doc.save(str(dst)), out)


# ODS writing: plain OpenDocument XML in a zip (no extra library)
_ODS_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
           'xmlns:number="urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0" '
           'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
           'office:version="1.2"')
_ODS_STYLES = (
    '<office:automatic-styles>'
    '<number:date-style style:name="N1"><number:year number:style="long"/><number:text>-</number:text>'
    '<number:month number:style="long"/><number:text>-</number:text><number:day number:style="long"/>'
    '</number:date-style>'
    '<number:date-style style:name="N2"><number:year number:style="long"/><number:text>-</number:text>'
    '<number:month number:style="long"/><number:text>-</number:text><number:day number:style="long"/>'
    '<number:text> </number:text><number:hours number:style="long"/><number:text>:</number:text>'
    '<number:minutes number:style="long"/><number:text>:</number:text><number:seconds number:style="long"/>'
    '</number:date-style>'
    '<number:time-style style:name="N3" number:truncate-on-overflow="false"><number:hours/>'
    '<number:text>:</number:text><number:minutes number:style="long"/><number:text>:</number:text>'
    '<number:seconds number:style="long"/></number:time-style>'
    '<style:style style:name="ceH" style:family="table-cell"><style:text-properties fo:font-weight="bold"/></style:style>'
    '<style:style style:name="ceD" style:family="table-cell" style:data-style-name="N1"/>'
    '<style:style style:name="ceDT" style:family="table-cell" style:data-style-name="N2"/>'
    '<style:style style:name="ceT" style:family="table-cell" style:data-style-name="N3"/>'
    '</office:automatic-styles>')


def _odf_text(s: str) -> str:
    """One <text:p> per line; runs of spaces and tabs encoded the ODF way (XML collapses them)."""
    paras = []
    for line in _clean(s).replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        out = []
        for m in re.finditer(r" +|\t|[^ \t]+", line):
            t = m.group(0)
            if t == "\t":
                out.append("<text:tab/>")
            elif t[0] != " ":
                out.append(html.escape(t, quote=False))
            elif m.start() == 0 or m.end() == len(line):
                out.append(f'<text:s text:c="{len(t)}"/>')
            else:
                out.append(" " + (f'<text:s text:c="{len(t) - 1}"/>' if len(t) > 1 else ""))
        paras.append("<text:p>" + "".join(out) + "</text:p>")
    return "".join(paras)


def _ods_cell(v, header: bool) -> str:
    st = ' table:style-name="ceH"' if header else ""
    if v is None:
        return "<table:table-cell/>"
    if isinstance(v, bool):
        return (f'<table:table-cell{st} office:value-type="boolean" office:boolean-value='
                f'"{str(v).lower()}"><text:p>{"TRUE" if v else "FALSE"}</text:p></table:table-cell>')
    if _is_number(v) and math.isfinite(v):
        t = cell_text(v)
        return f'<table:table-cell{st} office:value-type="float" office:value="{t}"><text:p>{t}</text:p></table:table-cell>'
    if isinstance(v, (dt.date, dt.datetime)):
        v2 = v.replace(tzinfo=None) if isinstance(v, dt.datetime) else v
        has_time = isinstance(v2, dt.datetime) and v2.time() != dt.time(0)
        iso = v2.isoformat(timespec="seconds") if has_time else (v2.date() if isinstance(v2, dt.datetime) else v2).isoformat()
        st = st or (' table:style-name="ceDT"' if has_time else ' table:style-name="ceD"')
        return (f'<table:table-cell{st} office:value-type="date" office:date-value="{iso}">'
                f'<text:p>{cell_text(v2)}</text:p></table:table-cell>')
    if isinstance(v, (dt.time, dt.timedelta)):
        secs = (v.total_seconds() if isinstance(v, dt.timedelta)
                else v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6)
        sign = "-" if secs < 0 else ""
        secs = abs(secs)
        h, rem = divmod(secs, 3600)
        m, s = divmod(rem, 60)
        st = st or ' table:style-name="ceT"'
        return (f'<table:table-cell{st} office:value-type="time" office:time-value='
                f'"{sign}PT{int(h)}H{int(m)}M{s:g}S"><text:p>{cell_text(v)}</text:p></table:table-cell>')
    return f'<table:table-cell{st} office:value-type="string">{_odf_text(cell_text(v))}</table:table-cell>'


def write_ods(tables: list[Table], out: str | Path) -> Path:
    out = Path(out)
    names = _sheet_names([n for n, _ in tables])

    def write(dst):
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet",
                        compress_type=zipfile.ZIP_STORED)
            zf.writestr("META-INF/manifest.xml", (
                '<?xml version="1.0" encoding="UTF-8"?>\n<manifest:manifest xmlns:manifest='
                '"urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                '<manifest:file-entry manifest:full-path="/" manifest:version="1.2" '
                'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>'
                '</manifest:manifest>'))
            zf.writestr("styles.xml", f'<?xml version="1.0" encoding="UTF-8"?>\n<office:document-styles '
                                      f'{_ODS_NS}><office:styles/></office:document-styles>')
            zf.writestr("meta.xml", '<?xml version="1.0" encoding="UTF-8"?>\n<office:document-meta '
                        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                        'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" office:version="1.2">'
                        '<office:meta><meta:generator>DropDial</meta:generator></office:meta>'
                        '</office:document-meta>')
            with zf.open("content.xml", "w") as f:
                w = io.TextIOWrapper(f, encoding="utf-8")
                w.write(f'<?xml version="1.0" encoding="UTF-8"?>\n<office:document-content {_ODS_NS}>'
                        f'{_ODS_STYLES}<office:body><office:spreadsheet>')
                for name, (_, rows) in zip(names, tables):
                    ncols = max((len(r) for r in rows), default=1) or 1
                    w.write(f'<table:table table:name="{html.escape(name)}">'
                            f'<table:table-column table:number-columns-repeated="{ncols}"/>')
                    for r, row in enumerate(rows):
                        w.write("<table:table-row>" + "".join(_ods_cell(v, r == 0) for v in row)
                                + "</table:table-row>")
                    if not rows:
                        w.write("<table:table-row><table:table-cell/></table:table-row>")
                    w.write("</table:table>")
                w.write("</office:spreadsheet></office:body></office:document-content>")
                w.flush()
                w.detach()
    _save(write, out)
    return out


# XLSM -> XLSX: drop the macro parts from the package; everything else is kept byte for byte
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
_MAIN_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


def _flat_xml(root, ns: str) -> bytes:
    """Serialize a one-level package part ([Content_Types].xml, *.rels) with its default namespace."""
    def name(tag):
        return tag.rsplit("}", 1)[-1]
    kids = "".join(f"<{name(e.tag)}" + "".join(f' {k}="{html.escape(v)}"' for k, v in e.attrib.items()) + "/>"
                   for e in root)
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<{name(root.tag)} xmlns="{ns}">{kids}</{name(root.tag)}>').encode()


def _strip_macros(src: Path, out: Path) -> bool:
    """False when the workbook can't be a plain .xlsx (it has Excel 4 macro sheets)."""
    from defusedxml.ElementTree import fromstring

    with _check_zip(src) as zf:
        names = zf.namelist()
        ct = fromstring(zf.read("[Content_Types].xml"))
        if any("macrosheet" in (e.get("ContentType") or "") or "dialogsheet" in (e.get("ContentType") or "")
               for e in ct):
            return False
        drop = {n for n in names if re.fullmatch(r"xl/(vbaProject|vbaProjectSignature\w*|vbaData)\.(bin|xml)", n)
                or n.startswith("xl/_rels/vbaProject")}
        for e in list(ct):
            ctype = e.get("ContentType") or ""
            part = (e.get("PartName") or "").lstrip("/")
            if part in drop or (e.tag == f"{{{_CT}}}Override" and "vbaProject" in ctype):
                ct.remove(e)
            elif ctype == "application/vnd.ms-excel.sheet.macroEnabled.main+xml":
                e.set("ContentType", _MAIN_XLSX)
            elif e.get("Extension", "").lower() == "bin" and "vbaProject" in ctype:
                if any(n.lower().endswith(".bin") and n not in drop for n in names):
                    e.set("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.printerSettings")
                else:
                    ct.remove(e)
        rels_name = "xl/_rels/workbook.xml.rels"
        rels = fromstring(zf.read(rels_name)) if rels_name in names else None
        if rels is not None:
            for e in list(rels):
                if "vbaProject" in (e.get("Type") or "") or "vbaProject" in (e.get("Target") or ""):
                    rels.remove(e)
        new = {"[Content_Types].xml": _flat_xml(ct, _CT)}
        if rels is not None:
            new[rels_name] = _flat_xml(rels, _RELS)

        def write(dst):
            with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
                for info in zf.infolist():
                    if info.filename in drop:
                        continue
                    data = new.get(info.filename)
                    zo.writestr(info.filename, data if data is not None else zf.read(info))
        _save(write, out)
    return True


# ---------------------------------------------------------------- family
ALL_TARGETS = ["XLSX", "CSV", "PDF", "JSON", "HTML", "DOCX", "MD", "ODS", "TSV"]
_OWN = {"xlsx": "XLSX", "csv": "CSV", "tsv": "TSV", "ods": "ODS"}

HINTS = {
    "XLSX": "Opens in Excel", "CSV": "Opens in Excel", "TSV": "Tab-separated text",
    "JSON": "Rows as objects", "HTML": "Web page table", "PDF": "Printable table",
    "DOCX": "Word table", "MD": "Markdown table", "ODS": "LibreOffice Calc", "VCF": "Contact cards",
    "Combine": "Files → one workbook", "Split sheets": "One XLSX per sheet",
}


def targets(src_ext: str) -> list[str]:
    ext = src_ext.lower().lstrip(".")
    out = [t for t in ALL_TARGETS if t != _OWN.get(ext)]
    if ext == "csv":
        out.append("VCF")       # a contact list → vCard (handled by the data family)
    return out


def convert(src: Path, target: str) -> list[Path]:
    src, target = Path(src), target.upper()
    ext = src.suffix.lower()
    if target == "VCF":
        from .data import csv_to_vcards
        return csv_to_vcards(src)
    if target == "XLSX" and ext == ".xlsm":
        out = sibling(src, ext=".xlsx")
        if _strip_macros(src, out):
            return [out]
    tables = _tables(src, typed=target in ("ODS", "JSON"))
    if target == "XLSX":
        return [write_xlsx(tables, sibling(src, ext=".xlsx"), infer=_is_text(src))]
    if target == "ODS":
        return [write_ods(tables, sibling(src, ext=".ods"))]
    if target in ("CSV", "TSV", "JSON"):
        written = []
        for out, (_, rows) in _per_sheet_outputs(src, tables, "." + target.lower()):
            if target == "JSON":
                _write_json(rows, out)
            else:
                _write_delimited(rows, out, "," if target == "CSV" else "\t", bom=target == "CSV")
            written.append(out)
        return written
    out = sibling(src, ext="." + target.lower())
    if target == "PDF":
        if sum(len(r) * len(r[0]) for _, r in tables) > MAX_PDF_CELLS:
            raise UserError("This table is too big for a PDF. Try XLSX or CSV instead.")
        _write_pdf(tables, out)
    elif target == "DOCX":
        if sum(len(r) * len(r[0]) for _, r in tables) > MAX_DOCX_CELLS:
            raise UserError("This table is too big for a Word document. Try XLSX or PDF instead.")
        _write_docx(tables, out)
    elif target == "HTML":
        _write_html(tables, src.stem, out)
    elif target == "MD":
        _write_md(tables, out)
    else:
        raise UserError("Can't convert a spreadsheet to {target}.", target=target)
    return [out]


def combine(files: list[Path]) -> list[Path]:
    """Several CSV/XLSX/… files -> one workbook with a sheet for each file (and each sheet)."""
    tables, infer = [], []
    for f in map(Path, files):
        try:
            ts = [(n, r) for n, r in read_tables(f) if not _empty(r)]   # empty files are skipped
        except UserError as e:
            raise UserError(f"{f.name}: {e}") from e
        tables += [(f.stem if len(ts) == 1 else f"{f.stem} - {n}", r) for n, r in ts]
        infer += [_is_text(f)] * len(ts)
    if not tables:
        raise UserError("These files are empty.")
    return [write_xlsx(tables, sibling(files[0], "Combined", ".xlsx"), infer=infer)]


def split_sheets(files: list[Path]) -> list[Path]:
    written = []
    for f in files:
        f = Path(f)
        tables = _tables(f, typed=False)
        if len(tables) < 2:
            raise UserError("This file has only one sheet with data.")
        for out, table in _per_sheet_outputs(f, tables, ".xlsx"):
            written.append(write_xlsx([table], out, infer=_is_text(f)))
    return written


register(Family(
    category="sheet",
    label="Spreadsheet",
    exts=EXTS,
    targets=targets,
    convert=convert,
    hints=HINTS,
    tools=[Tool("Combine", "⧉", HINTS["Combine"], combine, multi=True, working="Combining…"),
           Tool("Split sheets", "✂", HINTS["Split sheets"], split_sheets, working="Splitting sheets…")],
    available=_openpyxl_ok,
))
