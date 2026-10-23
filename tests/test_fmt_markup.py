"""Markdown / HTML / MHTML / RST / e-mail family: every output is parsed back and checked."""
from __future__ import annotations

import io
import re
from email.message import EmailMessage
from email.policy import default as email_policy
from html.parser import HTMLParser
from pathlib import Path

import pytest
from PIL import Image

from dropdial import convert as core, util
from dropdial.formats import markup
from dropdial.util import UserError

fitz = core.fitz

MD = """---
title: Front Matter Title
---
# Main Heading नमस्ते

Intro with **bold**, *italic*, `inline_code`, ~~struck~~ and a [link](https://example.com/page).

## Lists

- apple
- banana
    - nested cherry
- [x] done task

3. third
4. fourth

## Table

| Name | Qty |
|------|-----|
| Apple | 3 |
| 日本語 | 5 |

```python
def double(x):
    return x * 2
```

> A wise quote

![red box](pic.png)
![tracker](https://tracker.example/pixel.png)

<script>alert("x")</script>
<img src="x" onerror="alert(1)">
"""


def png_bytes(size=(120, 60), color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def md(tmp_path) -> Path:
    (tmp_path / "pic.png").write_bytes(png_bytes())
    p = tmp_path / "notes.md"
    p.write_text(MD, encoding="utf-8")
    return p


class Checker(HTMLParser):
    """Tag balance check: every non-void element is closed in order."""
    VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source",
            "track", "wbr"}

    def __init__(self):
        super().__init__()
        self.stack, self.errors, self.tags, self.attrs = [], [], [], []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attrs.append((tag, dict(attrs)))
        if tag not in self.VOID:
            self.stack.append(tag)

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1] != tag:
            self.errors.append(f"unexpected </{tag}> with stack {self.stack[-3:]}")
        else:
            self.stack.pop()


def check_html(text: str) -> Checker:
    c = Checker()
    c.feed(text)
    c.close()
    assert not c.errors and not c.stack, (c.errors, c.stack)
    return c


def pdf_text(p: Path) -> str:
    with fitz.open(p) as d:
        return "\n".join(pg.get_text() for pg in d)


def docx_doc(p: Path):
    from docx import Document
    return Document(p)


# ---------------------------------------------------------------- registry
def test_registry_and_targets():
    assert util.category("a.md") == "markup"
    assert util.category("a.HTML") == "markup"
    assert util.category("a.mhtml") == "markup"
    assert util.category("a.rst") == "markup"
    assert util.category("a.eml") == "email"
    assert core.targets_for("markup", "md") == ["PDF", "DOCX", "HTML", "TXT"]
    assert core.targets_for("markup", "html") == ["PDF", "DOCX", "MD", "TXT"]
    assert core.targets_for("email", "eml") == ["PDF", "HTML", "DOCX", "TXT"]
    from dropdial import formats
    for cat in ("markup", "email"):
        fam = formats.family(cat)
        for t in fam.targets("md") + fam.targets("eml") + [x.name for x in fam.tools]:
            assert len(fam.hints[t]) <= 22
    for ext, ts in markup.TARGETS.items():
        assert ext.upper() not in ts and len(ts) <= 9


# ---------------------------------------------------------------- Markdown
def test_md_to_pdf(md):
    [out] = core.convert(md, "markup", "PDF")
    assert out.suffix == ".pdf" and out.parent == md.parent
    text = pdf_text(out)
    for s in ("Main Heading", "bold", "inline_code", "apple", "nested cherry", "third", "fourth",
              "Apple", "日本語", "def double(x):", "return x * 2", "A wise quote", "☑"):
        assert s in text, s
    assert "alert" not in text and "Front Matter" not in text and "title:" not in text
    with fitz.open(out) as d:
        assert d.metadata["title"] == "Front Matter Title"
        toc = d.get_toc()
        assert [t[:2] for t in toc] == [[1, "Main Heading नमस्ते"], [2, "Lists"], [2, "Table"]]
        page = d[0]
        assert [l["uri"] for l in page.get_links()] == ["https://example.com/page"]
        imgs = page.get_images(full=True)
        assert len(imgs) == 1                             # local image in, remote tracker out
        pix = fitz.Pixmap(d, imgs[0][0])
        assert (pix.width, pix.height) == (120, 60)
        fonts = " ".join(f[3] for f in page.get_fonts())
        assert "Devanagari" in fonts and "Mono" in fonts  # Hindi shaped, code in monospace
    assert out.stat().st_size < 1_500_000                  # fonts are subset


def test_md_to_docx_structure(md):
    [out] = core.convert(md, "markup", "DOCX")
    d = docx_doc(out)
    paras = [(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]
    assert ("Heading 1", "Main Heading नमस्ते") in paras
    assert ("Heading 2", "Lists") in paras and ("Heading 2", "Table") in paras
    assert ("List Bullet", "apple") in paras and ("List Bullet 2", "nested cherry") in paras
    assert ("List Number", "third") in paras
    assert ("Quote", "A wise quote") in paras
    code = [p for p in d.paragraphs if p.style.name == "Code Block"]
    assert len(code) == 1 and "def double(x):\n    return x * 2" in code[0].text
    intro = next(p for p in d.paragraphs if p.text.startswith("Intro"))
    runs = {r.text: r for r in intro.runs}
    assert runs["bold"].bold and runs["italic"].italic and runs["struck"].font.strike
    assert runs["inline_code"].style.name == "Inline Code"
    assert "alert" not in "\n".join(p.text for p in d.paragraphs)
    # hyperlink with its external relationship
    links = [r for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert [r.target_ref for r in links] == ["https://example.com/page"]
    assert "link" in intro.text
    # the ordered list restarts at 3
    third = next(p for p in d.paragraphs if p.text == "third")
    num_id = third._p.pPr.numPr.numId.val
    num = d.part.numbering_part.element.num_having_numId(num_id)
    assert num.xpath("./w:lvlOverride/w:startOverride/@w:val") == ["3"]
    # table: real cells, header bold
    [t] = d.tables
    assert [[c.text for c in row.cells] for row in t.rows] == [["Name", "Qty"], ["Apple", "3"], ["日本語", "5"]]
    assert t.rows[0].cells[0].paragraphs[0].runs[0].bold
    # the picture is embedded once
    assert len(d.inline_shapes) == 1
    assert d.core_properties.title == "Front Matter Title"


def test_md_to_html(md):
    [out] = core.convert(md, "markup", "HTML")
    text = out.read_text(encoding="utf-8")
    c = check_html(text)
    assert text.startswith("<!DOCTYPE html>") and '<meta charset="utf-8">' in text
    assert "Content-Security-Policy" in text and "default-src 'none'" in text
    assert "script" not in c.tags and "onerror" not in text and "alert" not in text
    assert {"h1", "h2", "ul", "ol", "table", "pre", "blockquote", "a", "img"} <= set(c.tags)
    srcs = [a["src"] for t, a in c.attrs if t == "img"]
    assert len(srcs) == 1 and srcs[0].startswith("data:image/png;base64,")
    assert "tracker.example" not in text
    assert "<title>Front Matter Title</title>" in text and "नमस्ते" in text


def test_md_to_txt(md):
    [out] = core.convert(md, "markup", "TXT")
    text = out.read_text(encoding="utf-8")
    lines = text.splitlines()
    assert lines[0] == "Main Heading नमस्ते"
    assert "Intro with bold, italic, inline_code, struck and a link (https://example.com/page)." in lines
    assert "• apple" in lines and "  • nested cherry" in lines and "3. third" in lines
    assert "Apple\t3" in lines and "    return x * 2" in lines
    assert not re.search(r"\*\*|`|^#|\|---|<\w", text, re.M)


def test_md_never_overwrites(md):
    a = core.convert(md, "markup", "TXT")[0]
    b = core.convert(md, "markup", "TXT")[0]
    assert a != b and a.exists() and b.exists() and md.read_text(encoding="utf-8") == MD


# ---------------------------------------------------------------- HTML
PAGE = """<!DOCTYPE html><html lang="fr"><head><meta charset="windows-1252"><title>Caf\xe9 page</title>
<style>body{color:red}</style><link rel="stylesheet" href="https://cdn.example/x.css"></head>
<body onload="steal()"><h1>Caf\xe9 \u201cmenu\u201d</h1>
<p>Hello <b>bold <i>both</i></b> <a href="javascript:alert(1)">bad</a> <a href="https://ok.example/">good</a></p>
<ol><li>first<ol><li>inner one</li></ol></li><li>second</li></ol>
<table><tr><th>K</th><th>V</th></tr><tr><td>a|b</td><td>1</td></tr></table>
<table role="presentation"><tr><td><p>layout cell text</p></td></tr></table>
<iframe src="https://evil.example"></iframe><form><input value="secret"></form>
<pre><code class="language-js">let x = `tick`;</code></pre>
<img src="https://remote.example/a.png" alt="remote alt"></body></html>"""


@pytest.fixture
def page(tmp_path) -> Path:
    p = tmp_path / "page.html"
    p.write_bytes(PAGE.encode("cp1252"))
    return p


def test_html_to_md_roundtrip(page):
    import markdown
    [out] = core.convert(page, "markup", "MD")
    text = out.read_text(encoding="utf-8")
    assert text.startswith("# Café “menu”\n")
    assert "**bold** ***both***" in text
    assert "[good](https://ok.example/)" in text and "javascript" not in text and "bad" in text
    assert "1. first\n    1. inner one\n2. second" in text
    assert "| K | V |" in text and "a\\|b" in text
    assert "```js\nlet x = `tick`;\n```" in text
    assert "layout cell text" in text and "secret" not in text and "steal" not in text
    # the Markdown parses back into the same structure
    back = markdown.markdown(text, extensions=["tables", "fenced_code"])
    c = check_html(f"<div>{back}</div>")
    assert c.tags.count("li") == 3 and c.tags.count("th") == 2 and c.tags.count("td") == 2
    assert "<code class=\"language-js\">" in back


def test_html_to_docx_and_txt(page):
    [dx] = core.convert(page, "markup", "DOCX")
    d = docx_doc(dx)
    assert d.paragraphs[0].style.name == "Heading 1" and d.paragraphs[0].text == "Café “menu”"
    assert [(p.style.name, p.text) for p in d.paragraphs if "inner" in p.text] == [("List Number 2", "inner one")]
    assert len(d.tables) == 1                                   # layout table unwrapped
    assert d.tables[0].cell(1, 0).text == "a|b"
    [tx] = core.convert(page, "markup", "TXT")
    t = tx.read_text(encoding="utf-8")
    assert "Café “menu”" in t and "good (https://ok.example/)" in t and "bad (" not in t
    assert "1. first\n  1. inner one\n2. second" in t


def test_html_pdf_has_no_script_and_text(page):
    [out] = core.convert(page, "markup", "PDF")
    text = pdf_text(out)
    assert "Café" in text and "layout cell text" in text and "let x = `tick`;" in text
    assert "steal" not in text and "secret" not in text and "body{color" not in text


def test_bad_html_inputs(tmp_path):
    empty = tmp_path / "empty.html"
    empty.write_text("   \n", encoding="utf-8")
    with pytest.raises(UserError):
        core.convert(empty, "markup", "PDF")
    blank = tmp_path / "blank.html"
    blank.write_text("<html><body><script>x()</script></body></html>", encoding="utf-8")
    with pytest.raises(UserError):
        core.convert(blank, "markup", "DOCX")
    deep = tmp_path / "deep.html"
    deep.write_text("<div>" * 3000 + "deep text" + "</div>" * 3000, encoding="utf-8")
    with pytest.raises(UserError, match="nested too deeply"):  # never silently lose text
        core.convert(deep, "markup", "TXT")
    nested = tmp_path / "nested.html"                          # deep but parseable: flattened
    nested.write_text("<span>" * 200 + "inner text" + "</span>" * 200, encoding="utf-8")
    [out] = core.convert(nested, "markup", "TXT")
    assert out.read_text(encoding="utf-8").strip() == "inner text"
    with pytest.raises(UserError):
        core.convert(nested, "markup", "HTML")                 # not a target for .html


def test_local_image_rules(tmp_path):
    (tmp_path / "img").mkdir()
    (tmp_path / "img" / "ok.jpg").write_bytes(png_bytes())      # wrong extension, real PNG
    (tmp_path / "fake.png").write_bytes(b"not an image")
    (tmp_path / "pic.webp").write_bytes(b"")
    im = Image.new("RGB", (30, 20), (0, 0, 255))
    buf = io.BytesIO()
    im.save(buf, "WEBP")
    (tmp_path / "blue.webp").write_bytes(buf.getvalue())
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="40" height="20"><rect width="40" height="20" fill="green"/></svg>'
    (tmp_path / "g.svg").write_text(svg, encoding="utf-8")
    src = tmp_path / "imgs.md"
    src.write_text("![a](img/ok.jpg) ![b](fake.png) ![c](/etc/hostname) ![d](blue.webp) ![e](g.svg) "
                   "![f](file:///etc/passwd)\n", encoding="utf-8")
    [out] = core.convert(src, "markup", "DOCX")
    assert len(docx_doc(out).inline_shapes) == 3                 # ok.jpg, webp→png, svg→png
    [md_out] = core.convert(src, "markup", "HTML")
    c = check_html(md_out.read_text(encoding="utf-8"))
    srcs = [a["src"] for t, a in c.attrs if t == "img"]
    assert len(srcs) == 3 and all(s.startswith("data:image/png;base64,") for s in srcs)


# ---------------------------------------------------------------- MHTML
def make_mhtml(path: Path):
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.image import MIMEImage
    root = MIMEMultipart("related", type="text/html")
    root["Snapshot-Content-Location"] = "https://site.example/a/page.html"
    page = MIMEText('<html><head><title>Saved</title><link rel="stylesheet" href="style.css">'
                    '<script>evil()</script></head><body><h1>Saved page 東京</h1>'
                    '<p>Body text</p><img src="img/logo.png" alt="logo">'
                    '<img src="https://other.example/x.png" alt="gone"></body></html>', "html", "utf-8")
    page["Content-Location"] = "https://site.example/a/page.html"
    css = MIMEText("body{background:url(bg.png)} h1{color:#123456} @import url(https://x/y.css);",
                   "css", "utf-8")
    css["Content-Location"] = "https://site.example/a/style.css"
    img = MIMEImage(png_bytes((50, 40)), "png")
    img["Content-Location"] = "https://site.example/a/img/logo.png"
    bg = MIMEImage(png_bytes((4, 4), (0, 255, 0)), "png")
    bg["Content-Location"] = "https://site.example/a/bg.png"
    for p in (page, css, img, bg):
        root.attach(p)
    path.write_bytes(root.as_bytes())


def test_mhtml_to_html_pdf_txt(tmp_path):
    src = tmp_path / "saved.mhtml"
    make_mhtml(src)
    [out] = core.convert(src, "markup", "HTML")
    text = out.read_text(encoding="utf-8")
    c = check_html(text)
    assert "evil" not in text and "script" not in c.tags
    assert "Content-Security-Policy" in text
    srcs = [a["src"] for t, a in c.attrs if t == "img"]
    assert len(srcs) == 1 and srcs[0].startswith("data:image/png;base64,")
    import base64
    data = base64.b64decode(srcs[0].split(",", 1)[1])
    assert Image.open(io.BytesIO(data)).size == (50, 40)
    assert "h1{color:#123456}" in text                         # linked CSS inlined
    assert re.search(r'background:url\("data:image/png;base64,', text)
    assert "https://x/y.css" not in text and "other.example" not in text
    assert "東京" in text
    [pdf] = core.convert(src, "markup", "PDF")
    t = pdf_text(pdf)
    assert "Saved page 東京" in t and "Body text" in t and "evil" not in t
    with fitz.open(pdf) as d:
        assert len(d[0].get_images()) == 1
    [txt] = core.convert(src, "markup", "TXT")
    assert txt.read_text(encoding="utf-8").splitlines()[:3] == ["Saved page 東京", "", "Body text"]


def test_mhtml_garbage(tmp_path):
    src = tmp_path / "bad.mht"
    src.write_bytes(b"\x00\x01garbage" * 50)
    with pytest.raises(UserError):
        core.convert(src, "markup", "HTML")


# ---------------------------------------------------------------- e-mail
SUBJECT = "Réunion — 会議 🚀 update"


def make_eml(path: Path, html_body=True, attachments=True):
    msg = EmailMessage(policy=email_policy)
    msg["From"] = "Zoë Example <zoe@example.com>"
    msg["To"] = "Bob <bob@example.org>, ravi@example.in"
    msg["Cc"] = "Team <team@example.com>"
    msg["Date"] = "Tue, 03 Oct 2023 14:05:00 +0200"
    msg["Subject"] = SUBJECT
    msg.set_content("Plain version: hello Bob.\nSecond line https://example.com/x\n\nNew paragraph.")
    if html_body:
        msg.add_alternative('<html><body><p style="color:red">Hello <b>Bob</b> नमस्ते</p>'
                            '<img src="cid:logo123" alt="logo"><img src="https://track.example/p.gif">'
                            '<script>bad()</script><a href="javascript:x()">click</a></body></html>',
                            subtype="html")
        msg.get_payload()[1].add_related(png_bytes((64, 32)), "image", "png", cid="<logo123>",
                                         filename="logo.png", disposition="inline")
    if attachments:
        msg.add_attachment(b"evil content", maintype="text", subtype="plain", filename="../../evil.txt")
        msg.add_attachment(b"%PDF-1.4 fake", maintype="application", subtype="pdf", filename="résumé 2023.pdf")
        msg.add_attachment(b"second", maintype="text", subtype="plain", filename="..\\..\\EVIL.TXT")
        msg.add_attachment(b"con", maintype="application", subtype="octet-stream", filename="CON.txt")
        msg.add_attachment(b"noname", maintype="application", subtype="octet-stream", filename="..")
    path.write_bytes(msg.as_bytes())


def test_eml_raw_encoding_is_rfc2047(tmp_path):
    src = tmp_path / "m.eml"
    make_eml(src)
    raw = src.read_bytes()
    assert b"=?utf-8?" in raw.lower() and "会議".encode() not in raw.split(b"\n\n")[0]


def test_eml_to_pdf(tmp_path):
    src = tmp_path / "m.eml"
    make_eml(src)
    [out] = core.convert(src, "email", "PDF")
    t = pdf_text(out)
    for s in ("会議", "Zoë Example <zoe@example.com>", "Bob <bob@example.org>", "Team <team@example.com>",
              "14:05:00", "Hello", "Bob", "From", "Attachments", "evil.txt", "résumé 2023.pdf", "attachment 5.bin"):
        assert s in t, s
    assert "bad()" not in t and "Plain version" not in t           # HTML part preferred
    with fitz.open(out) as d:
        assert d.metadata["title"] == SUBJECT
        imgs = d[0].get_images(full=True)
        assert len(imgs) == 1                                    # cid image in, tracker out
        assert (imgs[0][2], imgs[0][3]) == (64, 32)


def test_eml_to_html_docx_txt(tmp_path):
    src = tmp_path / "m.eml"
    make_eml(src)
    [h] = core.convert(src, "email", "HTML")
    text = h.read_text(encoding="utf-8")
    c = check_html(text)
    assert f"<h1>{SUBJECT}</h1>" in text and "Zoë Example &lt;zoe@example.com&gt;" in text
    assert "track.example" not in text and "bad()" not in text and "javascript" not in text
    assert "style=" not in text.split("<main>")[1]
    assert sum(1 for t, a in c.attrs if t == "img" and a["src"].startswith("data:image/png")) == 1
    [dx] = core.convert(src, "email", "DOCX")
    d = docx_doc(dx)
    assert d.paragraphs[0].style.name == "Heading 1" and d.paragraphs[0].text == SUBJECT
    rows = {r.cells[0].text: r.cells[1].text for r in d.tables[0].rows}
    assert rows["From"] == "Zoë Example <zoe@example.com>" and rows["Cc"] == "Team <team@example.com>"
    assert any("नमस्ते" in p.text for p in d.paragraphs) and len(d.inline_shapes) == 1
    [tx] = core.convert(src, "email", "TXT")
    lines = tx.read_text(encoding="utf-8").splitlines()
    assert lines[0] == SUBJECT
    assert "From: Zoë Example <zoe@example.com>" in lines
    assert "Date: Tue, 03 Oct 2023 14:05:00 +0200" in lines
    assert any(l.startswith("Hello Bob नमस्ते") for l in lines)


def test_plain_text_eml(tmp_path):
    src = tmp_path / "plain.eml"
    make_eml(src, html_body=False, attachments=False)
    [h] = core.convert(src, "email", "HTML")
    c = check_html(h.read_text(encoding="utf-8"))
    assert ("a", {"href": "https://example.com/x"}) in c.attrs
    [tx] = core.convert(src, "email", "TXT")
    t = tx.read_text(encoding="utf-8")
    assert "Plain version: hello Bob.\nSecond line https://example.com/x\n\nNew paragraph." in t
    assert "Attachments" not in t
    with pytest.raises(UserError, match="no attachments"):
        markup.save_attachments(src)


def test_attachments_tool(tmp_path):
    src = tmp_path / "m.eml"
    make_eml(src)
    from dropdial import formats
    [tool] = formats.family("email").tools
    outs = tool.run([src])
    folder = tmp_path / "m attachments"
    assert all(p.parent == folder for p in outs)
    names = [p.name for p in outs]
    assert names == ["evil.txt", "résumé 2023.pdf", "EVIL (1).TXT", "_CON.txt", "attachment 5.bin"]
    assert [p.read_bytes() for p in outs] == [b"evil content", b"%PDF-1.4 fake", b"second", b"con", b"noname"]
    assert not (tmp_path.parent / "evil.txt").exists() and not (tmp_path / "evil.txt").exists()
    assert "logo.png" not in names                                # inline body image, not an attachment
    again = tool.run([src])                                       # second run: new folder, nothing overwritten
    assert again[0].parent.name == "m attachments (1)"
    md = tmp_path / "x.md"
    md.write_text("# hi", encoding="utf-8")
    with pytest.raises(UserError):
        markup.save_attachments(md)


def test_safe_name():
    assert markup._safe_name("../../etc/passwd", 1) == "passwd"
    assert markup._safe_name("C:\\Windows\\evil.exe", 1) == "evil.exe"
    assert markup._safe_name("a<b>:c?.txt", 1) == "a_b__c_.txt"
    assert markup._safe_name("nul", 1) == "_nul"
    assert markup._safe_name("", 3, "image/png") == "attachment 3.png"
    long = markup._safe_name("ü" * 300 + ".pdf", 1)
    assert long.endswith(".pdf") and len(long.encode()) <= 180


def test_not_an_email(tmp_path):
    src = tmp_path / "x.eml"
    src.write_bytes(b"just some text\nwithout headers\n")
    with pytest.raises(UserError, match="e-mail"):
        core.convert(src, "email", "PDF")


# ---------------------------------------------------------------- reStructuredText
def test_rst(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET", encoding="utf-8")
    src = tmp_path / "doc.rst"
    src.write_text(f"""Guide Title
===========

Some *emphasis* and ``literal`` text.

.. include:: {secret}

.. raw:: html

   <script>alert(1)</script>

- item one
- item two

=====  =====
A      B
=====  =====
x      y
=====  =====
""", encoding="utf-8")
    [pdf] = core.convert(src, "markup", "PDF")
    t = pdf_text(pdf)
    assert "Guide Title" in t and "item two" in t and "TOPSECRET" not in t and "alert" not in t
    [dx] = core.convert(src, "markup", "DOCX")
    d = docx_doc(dx)
    assert ("Heading 1", "Guide Title") in [(p.style.name, p.text) for p in d.paragraphs]
    lit = [r for p in d.paragraphs for r in p.runs if r.text == "literal"]
    assert lit and lit[0].style.name == "Inline Code"
    assert [[c.text for c in r.cells] for r in d.tables[0].rows] == [["A", "B"], ["x", "y"]]
    [md] = core.convert(src, "markup", "MD")
    assert md.read_text(encoding="utf-8").startswith("# Guide Title\n\nSome *emphasis* and `literal` text.")
