"""Structured data family (dropdial/formats/data.py): every output is re-read with an
independent parser (json, yaml, tomllib, plistlib, ElementTree, csv, openpyxl, PyMuPDF,
vobject) and compared value by value."""
import base64
import csv
import datetime as dt
import json
import plistlib
import subprocess
import sys
import tomllib
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
import yaml

from dropdial import convert as core, util
from dropdial.formats import data as D
from dropdial.util import UserError

UNI = {"hindi": "नमस्ते दुनिया", "japanese": "こんにちは世界", "emoji": "🎉😀"}

TYPICAL = {
    "name": "Ann Example",
    "age": 30,
    "height": 1.72,
    "active": True,
    "manager": None,
    "greetings": UNI,
    "tags": ["alpha", "beta", "gamma"],
    "projects": [
        {"@id": 7, "title": "Apollo", "budget": -1250.5, "done": False},
        {"@id": 8, "title": "हिन्दी प्रोजेक्ट", "budget": 99, "done": True},
    ],
    "note": {"@lang": "ja", "#text": "日本語のメモ"},
    "address": {"street": "1 Main St", "zip": "02134", "city": "Boston"},
}


def write(tmp_path, name, content):
    p = tmp_path / name
    if isinstance(content, bytes):
        p.write_bytes(content)
    else:
        p.write_text(content, encoding="utf-8")
    return p


def conv(src, target):
    out = D.convert(Path(src), target)
    assert len(out) == 1 and out[0].exists() and out[0] != Path(src)
    return out[0]


def jload(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


# ---------------------------------------------------------------- registry
def test_registry_claims_data_extensions_but_not_tables():
    for ext in ("json", "jsonl", "ndjson", "yaml", "yml", "toml", "xml", "plist", "vcf"):
        assert util.category(f"x.{ext}") == "data", ext
        t = core.targets_for("data", ext)
        assert 1 <= len(t) <= 9 and all(x == x.upper() for x in t)
        for name in t:
            assert name in D.HINTS and len(D.HINTS[name]) <= 22
    assert util.category("x.csv") != "data" and util.category("x.xlsx") != "data"
    assert "JSON" not in core.targets_for("data", "json")
    assert "YAML" not in core.targets_for("data", "yml")
    assert "PLIST" not in core.targets_for("data", "plist")
    assert core.targets_for("data", "vcf") == ["CSV", "XLSX", "JSON", "TXT"]


# ---------------------------------------------------------------- JSON <-> YAML / TOML
def test_json_to_yaml_exact(tmp_path):
    data = {**TYPICAL, "big": 123456789012345678901234567890}
    src = write(tmp_path, "a.json", json.dumps(data, ensure_ascii=False))
    out = conv(src, "YAML")
    text = out.read_text(encoding="utf-8")
    assert "नमस्ते" in text and "\\u" not in text      # real Unicode, not escapes
    assert yaml.safe_load(text) == data
    assert jload(conv(out, "JSON")) == data            # and back


def test_yaml_dates_become_iso_strings_and_stay_dates_in_toml(tmp_path):
    src = write(tmp_path, "c.yaml", "release: 2024-03-15\nstamp: 2024-03-15T10:20:30Z\n"
                                    "count: 3\nname: Zoë\n")
    j = jload(conv(src, "JSON"))
    assert j["release"] == "2024-03-15"
    assert dt.datetime.fromisoformat(j["stamp"]) == dt.datetime(2024, 3, 15, 10, 20, 30, tzinfo=dt.timezone.utc)
    assert j["count"] == 3 and j["name"] == "Zoë"
    t = tomllib.loads(conv(src, "TOML").read_text(encoding="utf-8"))
    assert t["release"] == dt.date(2024, 3, 15)
    assert t["stamp"] == dt.datetime(2024, 3, 15, 10, 20, 30, tzinfo=dt.timezone.utc)


def test_toml_to_json_and_yaml(tmp_path):
    src = write(tmp_path, "pyproject.toml", '[project]\nname = "demo"\nversion = "1.0"\n'
                '[server]\nport = 8080\nstarted = 1979-05-27T07:32:00\nwake = 07:30:00\n'
                'ratio = 0.5\n[[server.peers]]\nhost = "a"\n[[server.peers]]\nhost = "日本"\n')
    j = jload(conv(src, "JSON"))
    assert j == {"project": {"name": "demo", "version": "1.0"},
                 "server": {"port": 8080, "started": "1979-05-27T07:32:00", "wake": "07:30:00",
                            "ratio": 0.5, "peers": [{"host": "a"}, {"host": "日本"}]}}
    y = yaml.safe_load(conv(src, "YAML").read_text(encoding="utf-8"))
    assert y["server"]["started"] == dt.datetime(1979, 5, 27, 7, 32)
    assert y["server"]["wake"] == "07:30:00" and y["server"]["peers"][1]["host"] == "日本"


def test_json_to_toml_round_trip_and_conventions(tmp_path):
    data = {k: v for k, v in TYPICAL.items() if k != "manager"}
    src = write(tmp_path, "a.json", json.dumps(data, ensure_ascii=False))
    assert tomllib.loads(conv(src, "TOML").read_text(encoding="utf-8")) == data

    with pytest.raises(UserError, match="null.*manager"):
        D.convert(write(tmp_path, "n.json", json.dumps(TYPICAL)), "TOML")

    arr = [{"a": 1}, {"a": 2}]                          # TOML can't have a top-level array
    out = conv(write(tmp_path, "arr.json", json.dumps(arr)), "TOML")
    assert tomllib.loads(out.read_text(encoding="utf-8")) == {"items": arr}

    with pytest.raises(UserError, match="64 bits"):
        D.convert(write(tmp_path, "big.json", '{"n": 99999999999999999999}'), "TOML")


# ---------------------------------------------------------------- XML
def test_json_xml_json_round_trip(tmp_path):
    src = write(tmp_path, "a.json", json.dumps(TYPICAL, ensure_ascii=False))
    x = conv(src, "XML")
    root = ET.parse(x).getroot()                       # independent parser: well-formed, sane shape
    assert root.tag == "root"
    assert [p.get("id") for p in root.findall("projects")] == ["7", "8"]
    assert root.find("note").get("lang") == "ja" and root.find("note").text == "日本語のメモ"
    assert [t.text for t in root.findall("tags")] == ["alpha", "beta", "gamma"]
    assert root.find("greetings/emoji").text == "🎉😀"
    assert root.find("manager").text is None and len(root.find("manager")) == 0
    assert jload(conv(x, "JSON")) == TYPICAL


def test_xml_round_trip_top_level_list_and_single_root(tmp_path):
    arr = [{"id": 1, "v": "x"}, {"id": 2, "v": "ü"}]
    x = conv(write(tmp_path, "l.json", json.dumps(arr)), "XML")
    assert ET.parse(x).getroot().tag == "items"
    assert jload(conv(x, "JSON")) == arr

    one = {"catalog": {"@version": 2, "book": [{"title": "A"}, {"title": "B"}]}}
    x = conv(write(tmp_path, "c.json", json.dumps(one)), "XML")
    assert ET.parse(x).getroot().tag == "catalog"
    assert jload(conv(x, "JSON")) == one


XML_DOC = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library>
<lib:library xmlns:lib="urn:example:lib" xmlns="urn:example:default" open="true">
  <!-- comment -->
  <book id="b1" lang="hi">
    <title>गोदान</title>
    <year>1936</year>
    <price>1.50</price>
    <isbn>0012345</isbn>
  </book>
  <book id="b2">
    <title>吾輩は猫である</title>
    <year>1905</year>
    <price>2.5</price>
  </book>
  <lib:owner>Zoë <b>Smith</b> esq</lib:owner>
  <empty/>
</lib:library>
"""


def test_xml_to_json_convention(tmp_path):
    src = write(tmp_path, "lib.xml", XML_DOC)
    j = jload(conv(src, "JSON"))
    assert j == {"lib:library": {
        "@xmlns:lib": "urn:example:lib", "@xmlns": "urn:example:default", "@open": True,
        "book": [
            {"@id": "b1", "@lang": "hi", "title": "गोदान", "year": 1936, "price": "1.50", "isbn": "0012345"},
            {"@id": "b2", "title": "吾輩は猫である", "year": 1905, "price": 2.5},
        ],
        "lib:owner": {"#text": "Zoë esq", "b": "Smith"},
        "empty": None,
    }}
    # and the JSON goes back to equivalent XML (same prefixes, namespaces resolve the same)
    back = conv(src.with_name("lib.json"), "XML")
    r = ET.parse(back).getroot()
    assert r.tag == "{urn:example:lib}library"
    assert [b.findtext("{urn:example:default}title") for b in r.findall("{urn:example:default}book")] \
        == ["गोदान", "吾輩は猫である"]
    assert jload(conv(back, "JSON")) == j


def test_xml_entities_are_refused(tmp_path):
    secret = write(tmp_path, "secret.txt", "TOP-SECRET")
    laughs = ('<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
              '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]><lolz>&lol2;</lolz>')
    xxe = (f'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file://{secret}">]><r>&x;</r>')
    for i, doc in enumerate((laughs, xxe)):
        src = write(tmp_path, f"evil{i}.xml", doc)
        for target in ("JSON", "PDF"):
            with pytest.raises(UserError, match="entities"):
                D.convert(src, target)
        with pytest.raises(UserError, match="entities"):
            D.pretty([src])
    assert not list(tmp_path.glob("evil*.json"))


def test_xml_names_are_made_valid(tmp_path):
    src = write(tmp_path, "k.json", json.dumps({"first name": "A", "1st": 2, "ok": {"a/b": 3}}))
    x = conv(src, "XML")
    root = ET.parse(x).getroot()
    assert root.findtext("first_name") == "A" and root.findtext("_1st") == "2"
    assert root.find("ok").findtext("a_b") == "3"


# ---------------------------------------------------------------- JSON Lines
def test_jsonl_both_ways(tmp_path):
    rows = [{"id": i, "name": n} for i, n in enumerate(["a", "ब", "猫"])]
    src = write(tmp_path, "r.jsonl", "\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n\n")
    assert jload(conv(src, "JSON")) == rows
    back = conv(src.with_name("r.json"), "JSONL")
    lines = back.read_text(encoding="utf-8").splitlines()
    assert [json.loads(line) for line in lines] == rows
    # an object holding exactly one list (e.g. TOML's "items") is unwrapped
    wrapped = conv(write(tmp_path, "w.json", json.dumps({"items": rows})), "JSONL")
    assert [json.loads(x) for x in wrapped.read_text(encoding="utf-8").splitlines()] == rows

    bad = write(tmp_path, "bad.ndjson", '{"a": 1}\n{"a": \n')
    with pytest.raises(UserError, match="line 2"):
        D.convert(bad, "JSON")


# ---------------------------------------------------------------- property lists
def test_plist_both_encodings(tmp_path):
    native = {"Name": "Zoë 日本", "Count": 42, "Ratio": 0.25, "Enabled": True,
              "When": dt.datetime(2023, 7, 1, 12, 30), "Blob": b"\x00\x01\xffhi",
              "List": [1, "two", {"three": 3}]}
    src = write(tmp_path, "prefs.plist", plistlib.dumps(native, fmt=plistlib.FMT_BINARY))
    j = jload(conv(src, "JSON"))
    assert j["When"] == "2023-07-01T12:30:00"
    assert base64.b64decode(j["Blob"]) == native["Blob"]
    assert j["Name"] == native["Name"] and j["List"] == native["List"]

    x = conv(src, "XML")
    assert x.name == "prefs XML.plist" and x.read_bytes().startswith(b"<?xml")
    assert plistlib.loads(x.read_bytes()) == native
    b = conv(x, "BINARY")
    assert b.read_bytes().startswith(b"bplist00") and plistlib.loads(b.read_bytes()) == native
    assert src.read_bytes() == plistlib.dumps(native, fmt=plistlib.FMT_BINARY)   # untouched


def test_json_to_plist(tmp_path):
    data = {k: v for k, v in TYPICAL.items() if k != "manager"}
    out = conv(write(tmp_path, "a.json", json.dumps(data)), "PLIST")
    assert plistlib.loads(out.read_bytes()) == data
    with pytest.raises(UserError, match="null"):
        D.convert(write(tmp_path, "n.json", json.dumps(TYPICAL)), "PLIST")


def test_corrupt_plist(tmp_path):
    with pytest.raises(UserError, match="property list"):
        D.convert(write(tmp_path, "x.plist", b"bplist00garbage"), "JSON")
    evil = ('<?xml version="1.0"?><!DOCTYPE plist [<!ENTITY a "aaaa">]>'
            '<plist version="1.0"><string>&a;</string></plist>')
    with pytest.raises(UserError):
        D.convert(write(tmp_path, "e.plist", evil), "JSON")


# ---------------------------------------------------------------- tables
RECORDS = [
    {"id": 1, "name": "Ann", "address": {"city": "Pune", "geo": {"lat": 18.5}}, "tags": ["x", "y"],
     "formula": "=1+1", "big": 12345678901234567890},
    {"id": 2, "name": "नमस्ते", "address": {"city": "東京"}, "active": False, "note": None},
]


def test_csv_flattens_records(tmp_path):
    out = conv(write(tmp_path, "r.json", json.dumps(RECORDS, ensure_ascii=False)), "CSV")
    assert out.read_bytes().startswith(b"\xef\xbb\xbf")         # BOM so Excel reads UTF-8
    with open(out, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == ["id", "name", "address.city", "address.geo.lat", "tags", "formula",
                             "big", "active", "note"]
    assert rows[0]["address.geo.lat"] == "18.5" and json.loads(rows[0]["tags"]) == ["x", "y"]
    assert rows[0]["big"] == "12345678901234567890" and rows[0]["active"] == ""
    assert rows[1] == {"id": "2", "name": "नमस्ते", "address.city": "東京", "address.geo.lat": "",
                       "tags": "", "formula": "", "big": "", "active": "false", "note": ""}


def test_xlsx_types_and_no_formula_injection(tmp_path):
    from openpyxl import load_workbook
    out = conv(write(tmp_path, "r.yaml", yaml.safe_dump(RECORDS, allow_unicode=True, sort_keys=False)), "XLSX")
    ws = load_workbook(out).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0][:4] == ["id", "name", "address.city", "address.geo.lat"]
    assert rows[1][0] == 1 and rows[1][3] == 18.5 and rows[1][4] == '["x","y"]'
    f = ws.cell(2, 6)
    assert f.value == "=1+1" and f.data_type == "s"          # text, not a formula
    assert rows[1][6] == "12345678901234567890"             # too big for Excel's 15 digits: exact text
    assert rows[2][1] == "नमस्ते" and rows[2][2] == "東京" and rows[2][7] is False
    assert ws.freeze_panes == "A2" and ws.max_row == 3


def test_xml_to_table_and_dict_of_dicts(tmp_path):
    out = conv(write(tmp_path, "lib.xml", XML_DOC.replace("<lib:owner>Zoë <b>Smith</b> esq</lib:owner>\n  <empty/>", "")), "CSV")
    with open(out, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["title"] for r in rows] == ["गोदान", "吾輩は猫である"] and rows[0]["@id"] == "b1"

    users = {"alice": {"age": 30}, "bob": {"age": 41, "city": "Oslo"}}
    out = conv(write(tmp_path, "u.json", json.dumps(users)), "CSV")
    with open(out, encoding="utf-8-sig", newline="") as f:
        assert list(csv.reader(f)) == [["key", "age", "city"], ["alice", "30", ""], ["bob", "41", "Oslo"]]

    with pytest.raises(UserError, match="list of records"):
        D.convert(write(tmp_path, "one.json", '{"a": 1, "b": 2}'), "CSV")


# ---------------------------------------------------------------- PDF
def test_pdf_listing(tmp_path):
    import pymupdf
    long_token = "Z" * 300 + "END"
    data = {"intro": UNI, "long": long_token,
            "rows": [{"i": i, "v": f"value {i}"} for i in range(400)]}
    out = conv(write(tmp_path, "big.json", json.dumps(data, ensure_ascii=False)), "PDF")
    doc = pymupdf.open(out)
    assert doc.page_count >= 10
    text = "".join(p.get_text() for p in doc)
    flat = "".join(text.split())
    assert "こんにちは世界" in text
    assert long_token in flat                          # long lines wrap instead of being cut off
    assert '"v":"value399"' in flat                    # nothing missing at the end
    fonts = {f[3] for p in doc for f in p.get_fonts()}
    assert any("Devanagari" in f for f in fonts)       # Hindi drawn with a real Devanagari font
    assert out.stat().st_size < 1_000_000              # fallback fonts are subset


def test_pdf_from_yaml_keeps_comments_and_from_xml_is_pretty(tmp_path):
    import pymupdf
    src = write(tmp_path, "c.yml", "# deploy settings\nport: 80  # web\n")
    text = pymupdf.open(conv(src, "PDF"))[0].get_text()
    assert "# deploy settings" in text and "port: 80" in text
    src = write(tmp_path, "m.xml", "<a><b>1</b><c>二</c></a>")
    lines = pymupdf.open(conv(src, "PDF"))[0].get_text().splitlines()
    assert "  <b>1</b>" in lines and "  <c>二</c>" in lines


# ---------------------------------------------------------------- vCard
VCF = (
    "BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Ann Example\r\nN:Example;Ann;;;\r\n"
    "ORG:Acme\\, Inc.;Research\r\nTITLE:Chief Tinkerer\r\n"
    "EMAIL;TYPE=INTERNET,WORK:ann@acme.example\r\nEMAIL;TYPE=HOME:ann@home.example\r\n"
    "TEL;TYPE=CELL:+1 555 0100\r\nTEL;TYPE=WORK,VOICE:+1 555 0199\r\n"
    "ADR;TYPE=WORK:;;1 Main St;Springfield;IL;62701;USA\r\n"
    "BDAY:19900102\r\nURL:https://acme.example\r\n"
    "NOTE:Line one\\nLine two\\; with semicolon and a long tail that needs folding \r\n beyond 75 octets\r\n"
    "PHOTO;ENCODING=b;TYPE=JPEG:AAAA\r\n"
    "END:VCARD\r\n"
    "BEGIN:VCARD\r\nVERSION:2.1\r\nN;CHARSET=UTF-8;ENCODING=QUOTED-PRINTABLE:=E5=B1=B1=E7=94=B0;=E5=A4=AA=E9=83=8E\r\n"
    "TEL;HOME;VOICE:03-1234-5678\r\n"
    "NOTE;ENCODING=QUOTED-PRINTABLE:caf=C3=A9 =\r\nau lait\r\nEND:VCARD\r\n"
    "BEGIN:VCARD\r\nVERSION:4.0\r\nFN:प्रिया शर्मा\r\nitem1.EMAIL;PREF=1:priya@example.in\r\n"
    "item1.X-ABLabel:_$!<Home>!$_\r\nTEL;VALUE=uri;TYPE=\"cell,voice\":tel:+91-98765-43210\r\nEND:VCARD\r\n"
)


def test_vcard_parsing():
    c = D.parse_vcards(VCF)
    assert len(c) == 3
    ann, taro, priya = c
    assert ann["name"] == "Ann Example" and ann["given_name"] == "Ann" and ann["family_name"] == "Example"
    assert ann["organisation"] == "Acme, Inc., Research" and ann["title"] == "Chief Tinkerer"
    assert ann["emails"] == [{"value": "ann@acme.example", "types": ["work"]},
                             {"value": "ann@home.example", "types": ["home"]}]
    assert [p["value"] for p in ann["phones"]] == ["+1 555 0100", "+1 555 0199"]
    assert ann["addresses"][0]["locality"] == "Springfield" and ann["addresses"][0]["postal_code"] == "62701"
    assert ann["birthday"] == "1990-01-02" and ann["urls"] == ["https://acme.example"]
    assert ann["note"] == ("Line one\nLine two; with semicolon and a long tail that needs folding "
                           "beyond 75 octets")
    assert taro["name"] == "太郎 山田" and taro["family_name"] == "山田"
    assert taro["phones"] == [{"value": "03-1234-5678", "types": ["home", "voice"]}]
    assert taro["note"] == "café au lait"
    assert priya["name"] == "प्रिया शर्मा" and priya["emails"][0]["value"] == "priya@example.in"
    assert priya["phones"] == [{"value": "+91-98765-43210", "types": ["cell", "voice"]}]


def test_vcard_targets(tmp_path):
    from openpyxl import load_workbook
    src = write(tmp_path, "people.vcf", VCF)
    with open(conv(src, "CSV"), encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == D.CSV_COLUMNS and len(rows) == 3
    assert rows[0]["Email"] == "ann@acme.example; ann@home.example"
    assert rows[0]["Address"] == "1 Main St, Springfield IL 62701, USA"
    assert rows[0]["Note"].startswith("Line one\nLine two;")
    assert rows[1]["Name"] == "太郎 山田" and rows[2]["Phone"] == "+91-98765-43210"

    ws = load_workbook(conv(src, "XLSX")).active
    assert [c.value for c in ws[1]] == D.CSV_COLUMNS and ws.cell(4, 1).value == "प्रिया शर्मा"

    j = jload(conv(src, "JSON"))
    assert j == D.parse_vcards(VCF)

    txt = conv(src, "TXT").read_text(encoding="utf-8")
    assert "Ann Example\n  Organisation: Acme, Inc., Research" in txt
    assert "  Email:        ann@acme.example (work)" in txt and "café au lait" in txt


def test_csv_to_vcards_round_trip(tmp_path):
    vobject = pytest.importorskip("vobject")
    src = write(tmp_path, "people.vcf", VCF)
    table = conv(src, "CSV")
    out = D.csv_to_vcards(table)[0]
    raw = out.read_bytes()
    assert all(len(line) <= 75 for line in raw.split(b"\r\n"))       # folded per RFC
    again = D.parse_vcards(raw.decode("utf-8"))
    orig = D.parse_vcards(VCF)
    for a, b in zip(orig, again):
        assert a["name"] == b["name"] and a["organisation"] == b["organisation"]
        assert [e["value"] for e in a["emails"]] == [e["value"] for e in b["emails"]]
        assert [p["value"] for p in a["phones"]] == [p["value"] for p in b["phones"]]
        assert a["note"] == b["note"] and a["birthday"] == b["birthday"]
    cards = list(vobject.readComponents(raw.decode("utf-8")))          # independent parser
    assert [c.fn.value for c in cards] == ["Ann Example", "太郎 山田", "प्रिया शर्मा"]
    assert cards[0].note.value.startswith("Line one\nLine two;")
    assert [e.value for e in cards[0].contents["email"]] == ["ann@acme.example", "ann@home.example"]
    assert cards[0].org.value == ["Acme, Inc., Research"]

    # a hand-made Outlook-style CSV
    g = write(tmp_path, "g.csv", "First Name,Last Name,E-mail Address,Mobile Phone,Company\n"
                                 "Zoë,Ångström,zoe@example.com,+46 70 123,ACME\n,,,,\n")
    cards = list(vobject.readComponents(D.csv_to_vcards(g)[0].read_text(encoding="utf-8")))
    assert len(cards) == 1 and cards[0].fn.value == "Zoë Ångström"
    assert cards[0].n.value.family == "Ångström" and cards[0].tel.value == "+46 70 123"
    with pytest.raises(UserError, match="contact columns"):
        D.csv_to_vcards(write(tmp_path, "x.csv", "a,b\n1,2\n"))


def test_empty_vcard(tmp_path):
    with pytest.raises(UserError, match="contacts"):
        D.convert(write(tmp_path, "e.vcf", "hello\n"), "CSV")


# ---------------------------------------------------------------- tools
def test_pretty_and_minify_json(tmp_path):
    src = write(tmp_path, "a.json", json.dumps(TYPICAL, ensure_ascii=False))
    p = D.pretty([src])[0]
    m = D.minify([p])[0]
    assert p.name == "a Pretty.json" and m.name == "a Pretty Minified.json"
    assert json.loads(p.read_text(encoding="utf-8")) == TYPICAL
    assert p.read_text(encoding="utf-8").startswith('{\n  "name": "Ann Example",\n')
    mt = m.read_text(encoding="utf-8")
    assert json.loads(mt) == TYPICAL and "\n" not in mt and ": " not in mt


def test_pretty_and_minify_xml_keep_structure(tmp_path):
    src = write(tmp_path, "lib.xml", XML_DOC)
    p = D.pretty([src])[0]
    m = D.minify([src])[0]
    pt, mt = p.read_text(encoding="utf-8"), m.read_text(encoding="utf-8")
    assert "<lib:library" in pt and "<!-- comment -->" in pt and "\n    <title>गोदान</title>" in pt
    assert ">\n" not in mt.split("?>", 1)[1].strip() and "<title>गोदान</title>" in mt

    def canon(path):
        def walk(e):
            return (e.tag, sorted(e.attrib.items()), (e.text or "").strip(), [walk(c) for c in e])
        return walk(ET.parse(path).getroot())
    assert canon(p) == canon(src) == canon(m)


def test_pretty_plist_and_minify_jsonl(tmp_path):
    src = write(tmp_path, "x.plist", plistlib.dumps({"a": [1, 2], "b": "ü"}, fmt=plistlib.FMT_BINARY))
    p = D.pretty([src])[0]
    assert p.read_bytes().startswith(b"<?xml") and plistlib.loads(p.read_bytes()) == {"a": [1, 2], "b": "ü"}
    src = write(tmp_path, "r.jsonl", '{ "a" : 1 }\n{ "a" : 2 }\n')
    assert D.minify([src])[0].read_text() == '{"a":1}\n{"a":2}\n'
    with pytest.raises(UserError, match="Pretty works on"):
        D.pretty([write(tmp_path, "c.yaml", "a: 1\n")])


# ---------------------------------------------------------------- bad and hostile input
def test_bad_input_gives_friendly_errors(tmp_path):
    cases = [
        ("bad.json", '{"a": 1,,}', "line 1"),
        ("empty.json", "  \n", "empty"),
        ("bad.yaml", "a: [1, 2\n", "YAML"),
        ("bad.toml", "a = = 1", "TOML"),
        ("bad.xml", "<a><b></a>", "valid XML"),
        ("obj.yaml", "!!python/object/apply:os.system ['echo pwned']", "YAML"),
        ("deep.json", "[" * 100_000 + "]" * 100_000, "nested"),
        ("nan.yaml", "x: .nan\n", "NaN|nan"),
    ]
    for name, content, msg in cases:
        with pytest.raises(UserError, match=msg):
            D.convert(write(tmp_path, name, content), "YAML" if name.endswith(".json") else "JSON")


def test_yaml_alias_bomb_and_self_reference(tmp_path):
    bomb = "a: &a [x,x,x,x,x,x,x,x,x,x]\n" + "".join(
        f"{chr(98 + i)}: &{chr(98 + i)} [{', '.join(['*' + chr(97 + i)] * 10)}]\n" for i in range(8))
    with pytest.raises(UserError, match="too much data"):
        D.convert(write(tmp_path, "bomb.yaml", bomb), "JSON")
    with pytest.raises(UserError, match="refers to itself"):
        D.convert(write(tmp_path, "loop.yaml", "a: &a [*a]\n"), "JSON")


def test_size_guard_and_broken_text(tmp_path, monkeypatch):
    monkeypatch.setattr(D, "MAX_BYTES", 100)
    with pytest.raises(UserError, match="too big"):
        D.convert(write(tmp_path, "huge.json", json.dumps({"a": "x" * 200})), "YAML")
    monkeypatch.undo()
    with pytest.raises(UserError, match="broken text"):
        D.convert(write(tmp_path, "s.json", '[{"a": "\\ud800"}]'), "JSONL")
    with pytest.raises(UserError, match="control characters"):
        D.convert(write(tmp_path, "c.json", '{"a": "bell\\u0007"}'), "XML")


def test_outputs_never_overwrite(tmp_path):
    src = write(tmp_path, "a.yaml", "a: 1\n")
    first, second = conv(src, "JSON"), conv(src, "JSON")
    assert first.name == "a.json" and second.name == "a (1).json"
    assert src.read_text() == "a: 1\n" and jload(first) == jload(second) == {"a": 1}


def test_cli_goes_through_registry(tmp_path):
    src = write(tmp_path, "cfg.toml", 'title = "日本"\n[db]\nport = 5432\n')
    r = subprocess.run([sys.executable, "-m", "dropdial", "convert", str(src), "--to", "YAML"],
                       capture_output=True, text=True, cwd=Path(__file__).resolve().parents[1],
                       env={**__import__("os").environ, "QT_QPA_PLATFORM": "offscreen"})
    assert r.returncode == 0, r.stderr + r.stdout
    assert yaml.safe_load((tmp_path / "cfg.yaml").read_text(encoding="utf-8")) == \
        {"title": "日本", "db": {"port": 5432}}
