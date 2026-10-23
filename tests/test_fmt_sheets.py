"""Spreadsheet family: every output is parsed back with a different library or by hand."""
import csv
import datetime as dt
import io
import json
import re
import shutil
import unicodedata
import subprocess
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from dropdial import convert as core
from dropdial.formats import sheets
from dropdial.util import UserError, category

HINDI, JAPANESE, EMOJI = "नमस्ते दुनिया", "日本語のテキスト", "🎉"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
needs_lo = pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")


# ---------------------------------------------------------------- fixtures
def workbook(tmp_path, name="Book.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Name", "Qty", "Price", "Date", "Ok"])
    ws.append([f"Café {HINDI}", 3, 2.5, dt.date(2024, 3, 15), True])
    ws.append([f"{JAPANESE} {EMOJI}", 10, 0.1 + 0.2, dt.datetime(2024, 12, 1, 9, 30), False])
    ws.append(["Line one\nline two", -7, 1234567.891, dt.date(1999, 1, 31), None])
    ws["F1"] = "Total"
    ws["F2"] = "=B2*C2"            # openpyxl saves no cached value -> formula text is used
    staff = wb.create_sheet("Staff")
    staff.append(["Id", "Who"])
    staff.append(["00123", "Zoë"])
    wb.create_sheet("Empty")
    hidden = wb.create_sheet("Secret")
    hidden["A1"] = "hidden"
    hidden.sheet_state = "hidden"
    p = tmp_path / name
    wb.save(p)
    return p


def write_csv(tmp_path, text: str, name="data.csv", encoding="utf-8") -> Path:
    p = tmp_path / name
    p.write_bytes(text.encode(encoding))
    return p


def lo_convert(src: Path, fmt: str, outdir: Path) -> Path:
    outdir.mkdir(exist_ok=True)
    subprocess.run([SOFFICE, "--headless", f"-env:UserInstallation=file://{outdir}/profile",
                    "--convert-to", fmt, "--outdir", str(outdir), str(src)],
                   check=True, capture_output=True, timeout=180)
    return outdir / f"{src.stem}.{fmt.split(':')[0]}"


def read_csv_file(p: Path) -> list[list[str]]:
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


# ---------------------------------------------------------------- registry
def test_registered_and_targets():
    assert category("x.xlsx") == "sheet" and category("x.csv") == "sheet" and category("x.ods") == "sheet"
    for ext in ("xlsx", "xlsm", "csv", "tsv", "ods", "xls"):
        t = core.targets_for("sheet", ext)
        assert 0 < len(t) <= 9 and len(set(t)) == len(t)
        assert ext.upper() not in t
        assert all(t_ in sheets.HINTS and len(sheets.HINTS[t_]) <= 22 for t_ in t)
    assert core.targets_for("sheet", "csv")[0] == "XLSX"
    assert "XLSX" in core.targets_for("sheet", "xlsm")
    assert all(len(h) <= 22 for h in sheets.HINTS.values())


# ---------------------------------------------------------------- reading
def test_read_tables_values_and_formula_fallback(tmp_path):
    tables = sheets.read_tables(workbook(tmp_path))
    names = [n for n, _ in tables]
    assert names == ["Sales", "Staff", "Empty"]          # hidden sheet left out
    rows = dict(tables)["Sales"]
    assert rows[0] == ["Name", "Qty", "Price", "Date", "Ok", "Total"]
    assert rows[1][:5] == [f"Café {HINDI}", 3, 2.5, dt.datetime(2024, 3, 15), True]
    assert rows[1][5] == "=B2*C2"
    assert all(len(r) == 6 for r in rows)                # rectangular
    assert dict(tables)["Staff"][1] == ["00123", "Zoë"]


def test_lying_dimension_is_ignored(tmp_path):
    p = tmp_path / "dim.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.save(p)
    # rewrite <dimension> to claim the whole grid: rows must not be padded to 16384 columns
    with zipfile.ZipFile(p) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    files["xl/worksheets/sheet1.xml"] = re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:XFD1048576"',
                                               sheet).encode()
    with zipfile.ZipFile(p, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    assert sheets.read_tables(p) == [("Sheet", [["a", "b"]])]


@pytest.mark.parametrize("delim", [",", ";", "\t", "|"])
def test_csv_delimiter_sniffing(tmp_path, delim):
    rows = [["name", "city", "amount"], ["Ann", "Köln, DE", "1,5"], ["Bob", "Paris", "2"]]
    buf = io.StringIO()
    csv.writer(buf, delimiter=delim).writerows(rows)
    p = write_csv(tmp_path, buf.getvalue())
    assert sheets.read_tables(p)[0][1] == rows


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1252", "utf-16"])
def test_csv_encodings(tmp_path, encoding):
    p = write_csv(tmp_path, "name;note\nRenée;crème brûlée €5\n", encoding=encoding)
    assert sheets.read_tables(p)[0][1] == [["name", "note"], ["Renée", "crème brûlée €5"]]


def test_binary_garbage_is_a_friendly_error(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"\x00\x01garbage" * 100)
    with pytest.raises(UserError):
        sheets.convert(p, "CSV")
    z = tmp_path / "notabook.xlsx"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(UserError):
        sheets.convert(z, "CSV")
    e = write_csv(tmp_path, "\n\n", name="empty.csv")
    with pytest.raises(UserError, match="empty"):
        sheets.convert(e, "XLSX")


def test_zip_bomb_rejected(tmp_path, monkeypatch):
    p = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (3 << 20))
    monkeypatch.setattr(sheets, "MAX_UNZIPPED", 1 << 20)
    with pytest.raises(UserError, match="too big"):
        sheets.read_tables(p)


# ---------------------------------------------------------------- CSV / TSV / JSON
def test_xlsx_to_csv_one_file_per_sheet_with_bom_and_iso_dates(tmp_path):
    outs = sheets.convert(workbook(tmp_path), "CSV")
    assert [o.name for o in outs] == ["Book - Sales.csv", "Book - Staff.csv"]   # empty sheet skipped
    assert outs[0].read_bytes().startswith(b"\xef\xbb\xbf")
    rows = read_csv_file(outs[0])
    assert rows[1] == [f"Café {HINDI}", "3", "2.5", "2024-03-15", "TRUE", "=B2*C2"]
    assert rows[2] == [f"{JAPANESE} {EMOJI}", "10", "0.3", "2024-12-01 09:30:00", "FALSE", ""]
    assert rows[3][0] == "Line one\nline two" and rows[3][2] == "1234567.891"
    assert read_csv_file(outs[1]) == [["Id", "Who"], ["00123", "Zoë"]]


def test_single_sheet_keeps_plain_name_and_tsv(tmp_path):
    p = write_csv(tmp_path, 'a,b\n1,"x\ty"\n')
    out, = sheets.convert(p, "TSV")
    assert out.name == "data.tsv" and not out.read_bytes().startswith(b"\xef\xbb\xbf")
    with open(out, encoding="utf-8", newline="") as f:
        assert list(csv.reader(f, delimiter="\t")) == [["a", "b"], ["1", "x\ty"]]
    back, = sheets.convert(out, "CSV")
    assert read_csv_file(back) == [["a", "b"], ["1", "x\ty"]]


def test_json_rows_as_objects_with_types(tmp_path):
    outs = sheets.convert(workbook(tmp_path), "JSON")
    data = json.loads(outs[0].read_text(encoding="utf-8"))
    assert data[0] == {"Name": f"Café {HINDI}", "Qty": 3, "Price": 2.5, "Date": "2024-03-15", "Ok": True,
                       "Total": "=B2*C2"}
    assert data[1]["Price"] == 0.3 and data[1]["Date"] == "2024-12-01T09:30:00"
    assert data[2]["Ok"] is None and len(data) == 3
    assert json.loads(outs[1].read_text(encoding="utf-8")) == [{"Id": "00123", "Who": "Zoë"}]


def test_csv_to_json_infers_numbers_but_keeps_leading_zeros(tmp_path):
    p = write_csv(tmp_path, "zip,qty,price,when,,qty\n00123,4,3.50,2024-05-01,x,1\n,,-0.5,soon,,\n")
    out, = sheets.convert(p, "JSON")
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data == [
        {"zip": "00123", "qty": 4, "price": 3.5, "when": "2024-05-01", "Column 5": "x", "qty (2)": 1},
        {"zip": None, "qty": None, "price": -0.5, "when": "soon", "Column 5": None, "qty (2)": None},
    ]


# ---------------------------------------------------------------- XLSX
def test_csv_to_xlsx_types_text_and_no_formula_injection(tmp_path):
    p = write_csv(tmp_path, "zip,qty,price,date,cmd,big\n"
                            "00123,4,3.50,2024-05-01,=HYPERLINK(\"http://x\"),12345678901234567890\n"
                            f"7,-2,0.25,2024-13-45,{HINDI},1\n")
    out, = sheets.convert(p, "XLSX")
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.title == "data" and ws.freeze_panes == "A2" and ws["A1"].font.b
    assert ws["A2"].value == "00123" and ws["A2"].data_type == "s" and ws["A2"].number_format == "@"
    assert ws["B2"].value == 4 and ws["C2"].value == 3.5 and ws["C2"].number_format == "0.00"
    assert ws["D2"].value == dt.datetime(2024, 5, 1) and ws["D2"].is_date
    assert ws["E2"].value == '=HYPERLINK("http://x")' and ws["E2"].data_type == "s"   # text, not formula
    assert ws["F2"].value == "12345678901234567890"                  # too long for a number
    assert ws["D3"].value == "2024-13-45" and ws["E3"].value == HINDI and ws["B3"].value == -2
    with zipfile.ZipFile(out) as z:
        assert b"<f>" not in z.read("xl/worksheets/sheet1.xml")


def test_cp1252_csv_to_xlsx(tmp_path):
    p = write_csv(tmp_path, "Name,Price\nCafé Noël,€3\n", encoding="cp1252")
    out, = sheets.convert(p, "XLSX")
    assert [c.value for c in openpyxl.load_workbook(out).active[2]] == ["Café Noël", "€3"]


def test_xlsm_to_xlsx_strips_macros_keeps_everything_else(tmp_path):
    src = workbook(tmp_path, "Macro.xlsx")
    xlsm = tmp_path / "Macro.xlsm"
    with zipfile.ZipFile(src) as z:
        files = {n: z.read(n) for n in z.namelist()}
    ct = files["[Content_Types].xml"].decode()
    ct = ct.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
    ct = ct.replace("<Default Extension", '<Default Extension="bin" ContentType="application/vnd.ms-office.'
                    'vbaProject"/><Default Extension', 1)
    files["[Content_Types].xml"] = ct.encode()
    rels = files["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rIdVBA" Type="http://schemas.microsoft.com/office/2006/'
        'relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
    files["xl/_rels/workbook.xml.rels"] = rels.encode()
    files["xl/vbaProject.bin"] = b"\xd0\xcf\x11\xe0fake macro project"
    with zipfile.ZipFile(xlsm, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)

    assert "XLSX" in sheets.targets("xlsm")
    out, = sheets.convert(xlsm, "XLSX")
    assert out.suffix == ".xlsx" and xlsm.exists()
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        ct = z.read("[Content_Types].xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
    assert "xl/vbaProject.bin" not in names
    assert "macroEnabled" not in ct and "vbaProject" not in ct and "vbaProject" not in rels
    ET.fromstring(ct)
    wb = openpyxl.load_workbook(out)
    assert wb["Sales"]["F2"].value == "=B2*C2"          # formulas survive (it's the same package)
    assert wb["Staff"]["A2"].value == "00123"
    if SOFFICE:
        csv_out = lo_convert(out, "csv", tmp_path / "lo")
        assert read_csv_file(csv_out)[0][:2] == ["Name", "Qty"]


# ---------------------------------------------------------------- ODS
def _ods_rows(path: Path) -> dict[str, list[list[str]]]:
    """Hand-written check: the visible text of each cell."""
    t = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    x = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    o = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
    with zipfile.ZipFile(path) as z:
        assert z.namelist()[0] == "mimetype" and z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        assert z.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
        ET.fromstring(z.read("META-INF/manifest.xml"))
        root = ET.fromstring(z.read("content.xml"))
    out = {}
    for table in root.iter(t + "table"):
        rows = []
        for tr in table.iter(t + "table-row"):
            rows.append([(c.get(o + "value-type"), "\n".join("".join(p.itertext()) for p in c.findall(x + "p")))
                         for c in tr.findall(t + "table-cell")])
        out[table.get(t + "name")] = rows
    return out


def test_xlsx_to_ods_and_back(tmp_path):
    out, = sheets.convert(workbook(tmp_path), "ODS")
    got = _ods_rows(out)
    assert list(got) == ["Sales", "Staff"]
    assert got["Sales"][1][:5] == [("string", f"Café {HINDI}"), ("float", "3"), ("float", "2.5"),
                                   ("date", "2024-03-15"), ("boolean", "TRUE")]
    assert got["Sales"][3][0] == ("string", "Line one\nline two")
    assert got["Staff"][1][0] == ("string", "00123")
    # our ODS reader gets the same values back
    back = dict(sheets.read_tables(out))
    assert back["Sales"][1][:5] == [f"Café {HINDI}", 3, 2.5, dt.date(2024, 3, 15), True]
    assert back["Sales"][2][3] == dt.datetime(2024, 12, 1, 9, 30)
    assert back["Staff"] == [["Id", "Who"], ["00123", "Zoë"]]


def test_ods_spaces_tabs_roundtrip(tmp_path):
    tricky = ["  lead", "a  b   c", "tab\there", "trail  ", "x < y & z"]
    p = tmp_path / "t.ods"
    sheets.write_ods([("T", [["h"], *[[s] for s in tricky]])], p)
    assert [r[0] for r in sheets.read_tables(p)[0][1][1:]] == tricky


@needs_lo
def test_ods_opens_in_libreoffice(tmp_path):
    out, = sheets.convert(workbook(tmp_path), "ODS")
    rows = read_csv_file(lo_convert(out, "csv", tmp_path / "lo"))
    assert rows[0] == ["Name", "Qty", "Price", "Date", "Ok", "Total"]
    assert rows[1][:3] == [f"Café {HINDI}", "3", "2.5"] and rows[1][3] == "2024-03-15"


@needs_lo
def test_read_libreoffice_ods(tmp_path):
    ods = lo_convert(workbook(tmp_path), "ods", tmp_path / "lo")
    tables = dict(sheets.read_tables(ods))
    assert tables["Sales"][1][:5] == [f"Café {HINDI}", 3, 2.5, dt.date(2024, 3, 15), True]
    assert tables["Sales"][3][0] == "Line one\nline two"
    assert tables["Staff"][1] == ["00123", "Zoë"]
    out, = sheets.convert(ods, "XLSX")
    assert openpyxl.load_workbook(out)["Sales"]["C3"].value == 0.3


# ---------------------------------------------------------------- XLS (read only)
def test_xls_read(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Old")
    datefmt = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for c, v in enumerate(["Name", "Qty", "When"]):
        ws.write(0, c, v)
    ws.write(1, 0, f"Zoë {JAPANESE}")
    ws.write(1, 1, 42)
    ws.write(1, 2, dt.date(2020, 2, 29), datefmt)
    p = tmp_path / "old.xls"
    wb.save(str(p))
    assert category(p) == "sheet"
    assert sheets.read_tables(p) == [("Old", [["Name", "Qty", "When"],
                                              [f"Zoë {JAPANESE}", 42, dt.datetime(2020, 2, 29)]])]
    out, = sheets.convert(p, "XLSX")
    ws2 = openpyxl.load_workbook(out).active
    assert ws2["B2"].value == 42 and ws2["C2"].value == dt.datetime(2020, 2, 29)


def test_text_masquerading_as_xls(tmp_path):
    p = tmp_path / "export.xls"
    p.write_text("a\tb\n1\t002\n", encoding="utf-8")
    assert sheets.read_tables(p)[0][1] == [["a", "b"], ["1", "002"]]
    h = tmp_path / "page.xls"
    h.write_text("<html><table><tr><td>1</td></tr></table></html>")
    with pytest.raises(UserError, match="web page"):
        sheets.read_tables(h)


# ---------------------------------------------------------------- HTML / MD / DOCX / PDF
class _Cells(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self.heads, self._cell, self._tag = [], [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr":
            self.tables[-1].append([])
        elif tag in ("td", "th"):
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append("\n")
        elif tag == "h2":
            self._tag = "h2"

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self.tables[-1][-1].append("".join(self._cell))
            self._cell = None
        if tag == "h2":
            self._tag = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)
        elif self._tag == "h2":
            self.heads.append(data)


def test_html_table(tmp_path):
    out, = sheets.convert(workbook(tmp_path), "HTML")
    text = out.read_text(encoding="utf-8")
    assert text.startswith("<!doctype html>") and '<meta charset="utf-8">' in text
    p = _Cells()
    p.feed(text)
    assert p.heads == ["Sales", "Staff"]
    assert p.tables[0][0] == ["Name", "Qty", "Price", "Date", "Ok", "Total"]
    assert p.tables[0][1] == [f"Café {HINDI}", "3", "2.5", "2024-03-15", "TRUE", "=B2*C2"]
    assert p.tables[0][3][0] == "Line one\nline two"
    assert p.tables[1] == [["Id", "Who"], ["00123", "Zoë"]]
    evil = write_csv(tmp_path, "a\n<script>alert(1)</script>\n", name="evil.csv")
    out2, = sheets.convert(evil, "HTML")
    assert "<script>alert" not in out2.read_text(encoding="utf-8")


def _md_split(line: str) -> list[str]:
    cells = re.split(r"(?<!\\) \| ", line.strip()[2:-2])
    return [c.replace("\\|", "|").replace("\\\\", "\\").replace("<br>", "\n") for c in cells]


def test_markdown_table(tmp_path):
    p = write_csv(tmp_path, 'item,qty,note\npipe,1,"a|b"\nslash,22,"c:\\temp"\nmulti,3,"x\ny"\n')
    out, = sheets.convert(p, "MD")
    lines = out.read_text(encoding="utf-8").splitlines()
    assert _md_split(lines[0]) == ["item", "qty", "note"]
    assert lines[1] == "| --- | ---: | --- |"                  # numeric column right-aligned
    assert [_md_split(x) for x in lines[2:]] == [["pipe", "1", "a|b"], ["slash", "22", "c:\\temp"],
                                                 ["multi", "3", "x\ny"]]
    multi, = sheets.convert(workbook(tmp_path), "MD")
    assert "## Sales" in multi.read_text(encoding="utf-8") and "## Staff" in multi.read_text(encoding="utf-8")


def test_docx_table(tmp_path):
    import docx
    out, = sheets.convert(workbook(tmp_path), "DOCX")
    d = docx.Document(out)
    assert [p.text for p in d.paragraphs if p.style.name.startswith("Heading")] == ["Sales", "Staff"]
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["Name", "Qty", "Price", "Date", "Ok", "Total"]
    assert [c.text for c in t.rows[1].cells] == [f"Café {HINDI}", "3", "2.5", "2024-03-15", "TRUE", "=B2*C2"]
    assert t.rows[3].cells[0].text == "Line one\nline two"
    assert t.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert t._tbl.tr_lst[0].trPr.xml.count("tblHeader") == 1       # header repeats on each page
    assert [c.text for c in d.tables[1].rows[1].cells] == ["00123", "Zoë"]


def test_docx_wide_goes_landscape(tmp_path):
    import docx
    p = write_csv(tmp_path, ",".join(f"c{i}" for i in range(10)) + "\n" + ",".join("1" * 10) + "\n")
    out, = sheets.convert(p, "DOCX")
    sec = docx.Document(out).sections[0]
    assert sec.page_width > sec.page_height


def test_pdf_repeats_header_on_every_page(tmp_path):
    fitz = core.fitz
    rows = ["Code,Product,Amount"] + [f"R{i:04d},Item {i} {JAPANESE},{i}.25" for i in range(400)]
    p = write_csv(tmp_path, "\n".join(rows) + "\n")
    out, = sheets.convert(p, "PDF")
    with fitz.open(out) as doc:
        assert doc.page_count >= 5
        seen = []
        for page in doc:
            assert page.rect.width < page.rect.height                    # narrow table: portrait
            text = page.get_text()
            assert text.lstrip().startswith("Code\nProduct\nAmount")       # header on every page
            seen += re.findall(r"R\d{4}", text)
        assert seen == [f"R{i:04d}" for i in range(400)]                  # every row exactly once, in order
        assert JAPANESE in doc[0].get_text()


def test_pdf_wide_sheet_landscape_and_fits(tmp_path):
    fitz = core.fitz
    head = ",".join(f"Heading number {i}" for i in range(22))
    p = write_csv(tmp_path, head + "\n" + ",".join(str(i * 1000) for i in range(22)) + "\n")
    out, = sheets.convert(p, "PDF")
    with fitz.open(out) as doc:
        page = doc[0]
        assert page.rect.width > page.rect.height
        words = page.get_text("words")
        assert all(0 <= w[0] and w[2] <= page.rect.width for w in words)  # nothing runs off the page
        assert "21000" in [w[4] for w in words]


def test_pdf_multi_sheet_headings_and_hindi(tmp_path):
    fitz = core.fitz
    out, = sheets.convert(workbook(tmp_path), "PDF")
    with fitz.open(out) as doc:
        assert doc.page_count == 2                                         # one per sheet
        texts = [unicodedata.normalize("NFKC", pg.get_text()) for pg in doc]   # "ﬀ" ligature -> "ff"
        assert texts[0].startswith("Sales") and texts[1].startswith("Staff")
        assert "Café" in doc[0].get_text() and "00123" in doc[1].get_text()
        fonts = " ".join(f[3] for f in doc[0].get_fonts())
        assert "Devanagari" in fonts                                       # Hindi drawn with a real font


# ---------------------------------------------------------------- tools
def _tool(name):
    return next(t for t in sheets.convert.__globals__["register"].__globals__["FAMILIES"]["sheet"].tools
                if t.name == name)


def test_combine_tool(tmp_path):
    a = write_csv(tmp_path, "x,y\n1,00007\n", name="first.csv")
    b = workbook(tmp_path)
    c = write_csv(tmp_path, "", name="blank.csv")
    tool = _tool("Combine")
    assert tool.multi
    out, = tool.run([a, b, c])
    assert out.name == "first Combined.xlsx"
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["first", "Book - Sales", "Book - Staff"]
    assert wb["first"]["A2"].value == 1 and wb["first"]["B2"].value == "00007"
    assert wb["Book - Sales"]["B2"].value == 3 and wb["Book - Sales"]["F2"].value == "=B2*C2"
    assert wb["Book - Sales"]["F2"].data_type == "s"


def test_split_sheets_tool(tmp_path):
    outs = _tool("Split sheets").run([workbook(tmp_path)])
    assert [o.name for o in outs] == ["Book - Sales.xlsx", "Book - Staff.xlsx"]
    s = openpyxl.load_workbook(outs[0]).active
    assert s.title == "Sales" and s["A2"].value == f"Café {HINDI}" and s["D2"].value == dt.datetime(2024, 3, 15)
    with pytest.raises(UserError):
        _tool("Split sheets").run([write_csv(tmp_path, "a\n1\n")])


def test_never_overwrites(tmp_path):
    p = write_csv(tmp_path, "a\n1\n")
    first, = sheets.convert(p, "XLSX")
    second, = sheets.convert(p, "XLSX")
    assert first != second and first.exists() and second.name == "data (1).xlsx"


def test_sheet_names_are_valid():
    names = sheets._sheet_names(["a/b:c", "A" * 40, "a" * 40, "", "x[1]"])
    assert names[0] == "a_b_c" and all(len(n) <= 31 for n in names)
    assert len({n.lower() for n in names}) == 5


def test_write_xlsx_helper(tmp_path):
    out = sheets.write_xlsx([("T", [["k", "v"], ["a", 1.5], ["b", None]])], tmp_path / "h.xlsx")
    assert [[c.value for c in r] for r in openpyxl.load_workbook(out).active.iter_rows()] == \
        [["k", "v"], ["a", 1.5], ["b", None]]


def test_cli_convert(tmp_path):
    p = write_csv(tmp_path, "a,b\n1,2\n")
    r = subprocess.run(["python", "-m", "dropdial", "convert", str(p), "--to", "JSON"],
                       capture_output=True, text=True, cwd=Path(__file__).parent.parent, timeout=120)
    assert r.returncode == 0, r.stderr
    assert json.loads((tmp_path / "data.json").read_text()) == [{"a": 1, "b": 2}]
